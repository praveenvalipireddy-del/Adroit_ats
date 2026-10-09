"""Per-consultant job application tracker.

One `applications` row per consultant + job (the table the Reporting board already uses), with the
whole life of the application: status, how it was applied (portal / email / both), client and
recruiter contact, rate, which saved resume version was sent, follow-up date, interview date /
round, notes - and every change in `application_events` (who, when, what).

Statuses are stored in `applications.stage` (the existing board keys stay valid) and shown with
the recruiter's wording via STATUS_LABELS. Applications are created automatically by drafts
(models.mark_job_drafted) and the Jobs "Applied" dropdown, or by hand for applications made outside
the app. check_replies() reads the consultant's connected Gmail (read-only) for mail FROM each
open application's recruiter address and moves it to "Recruiter responded".
"""
import imaplib
import io
import re
from datetime import date, datetime, timedelta
from typing import Dict, List, Optional

STATUS_LABELS = {
    "Drafted": "Draft created", "Applied": "Applied", "Responded": "Recruiter responded", "Screening": "Recruiter responded",
    "Submitted": "Submitted to client", "Interviewing": "Interview", "Offer": "Offer", "Hired": "Placed",
    "Rejected": "Rejected", "No response": "No response", "On hold": "On hold", "Withdrawn": "Withdrawn", "Saved": "Saved",
}
STATUSES = ["Drafted", "Applied", "Responded", "Submitted", "Interviewing", "Offer", "Hired", "Rejected", "No response", "On hold", "Withdrawn"]
OPEN = {"Saved", "Drafted", "Applied", "Responded", "Screening", "Submitted", "Interviewing", "Offer"}
ORDER = {s: i for i, s in enumerate(["Saved", "Drafted", "Applied", "Responded", "Screening", "Submitted", "Interviewing", "Offer", "Hired"])}
METHODS = {"": "", "portal": "Applied in portal", "email": "Email sent", "both": "Portal + email"}
EDITABLE = ("apply_method", "client_company", "recruiter_name", "recruiter_email", "recruiter_phone", "rate", "follow_up_date",
            "interview_at", "interview_round", "notes")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(r"^\d{4}-\d{2}-\d{2}(T\d{2}:\d{2})?$")


class TrackerError(ValueError):
    pass


def event(cur, app_id: int, user_id: Optional[int], kind: str, detail: str):
    cur.execute("INSERT INTO application_events (application_id, user_id, kind, detail) VALUES (?, ?, ?, ?)",
                (app_id, user_id, kind, detail[:500]))


def upsert(conn, job_id: int, candidate_id: int, user_id: Optional[int], stage: Optional[str] = None, **fields) -> int:
    """The application for this consultant + job (created if missing). `stage` only moves it FORWARD
    along the normal path (a new draft never pulls an Interview back to 'Draft created')."""
    cur = conn.cursor()
    cur.execute("SELECT id, stage FROM applications WHERE job_id = ? AND candidate_id = ?", (job_id, candidate_id))
    row = cur.fetchone()
    if row:
        app_id, current = row[0], row[1] or "Saved"
    else:
        cur.execute("INSERT INTO applications (job_id, candidate_id, user_id, stage) VALUES (?, ?, ?, ?)",
                    (job_id, candidate_id, user_id, stage or "Saved"))
        app_id = cur.lastrowid
        if not app_id:
            cur.execute("SELECT id FROM applications WHERE job_id = ? AND candidate_id = ?", (job_id, candidate_id))
            app_id = cur.fetchone()[0]
        current = stage or "Saved"
        event(cur, app_id, user_id, "created", f"Tracked - {STATUS_LABELS.get(current, current)}")
    if stage and stage != current and ORDER.get(stage, 99) > ORDER.get(current, -1) and current in ORDER:
        cur.execute("UPDATE applications SET stage = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (stage, app_id))
        event(cur, app_id, user_id, "status", f"{STATUS_LABELS.get(current, current)} -> {STATUS_LABELS.get(stage, stage)}")
    fields = {k: v for k, v in fields.items() if v not in (None, "")}
    if fields:
        cur.execute(f"UPDATE applications SET {', '.join(k + ' = ?' for k in fields)}, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                    list(fields.values()) + [app_id])
    if stage == "Applied" and not (row and current not in ("Saved", "Drafted")):
        cur.execute("UPDATE applications SET applied_date = CURRENT_TIMESTAMP WHERE id = ?", (app_id,))
    conn.commit()
    return app_id


def update(conn, app_id: int, user: Dict, data: Dict) -> Dict:
    cur = conn.cursor()
    cur.execute("SELECT * FROM applications WHERE id = ?", (app_id,))
    row = cur.fetchone()
    if not row:
        raise LookupError("Application not found")
    row = dict(row)
    changes = {}
    if "status" in data:
        status = str(data["status"] or "")
        if status not in STATUSES:
            raise TrackerError("Unknown status.")
        if status != row["stage"]:
            changes["stage"] = status
            event(cur, app_id, user["id"], "status", f"{STATUS_LABELS.get(row['stage'], row['stage'])} -> {STATUS_LABELS[status]}")
            if status == "Applied" and row["stage"] in ("Saved", "Drafted", None):
                changes["applied_date"] = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
    for k in EDITABLE:
        if k not in data:
            continue
        v = str(data[k] or "").strip()
        if k == "apply_method" and v not in METHODS:
            raise TrackerError("Unknown apply method.")
        if k == "follow_up_date" and v and not _DATE_RE.match(v):
            raise TrackerError("Follow-up date must be YYYY-MM-DD.")
        if k == "interview_at" and v and not _DATETIME_RE.match(v):
            raise TrackerError("Interview date must be YYYY-MM-DD (optionally with a time).")
        if k == "recruiter_email" and v and "@" not in v:
            raise TrackerError("That recruiter email doesn't look right.")
        v = v[:2000] if k == "notes" else v[:200]
        if v != (row.get(k) or ""):
            changes[k] = v
            label = k.replace("_", " ")
            event(cur, app_id, user["id"], "note" if k == "notes" else "field",
                  "Notes updated" if k == "notes" else f"{label}: {v or '(cleared)'}")
    if changes:
        cur.execute(f"UPDATE applications SET {', '.join(k + ' = ?' for k in changes)}, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                    list(changes.values()) + [app_id])
    conn.commit()
    return get(conn, app_id)


_SELECT = """SELECT a.id, a.job_id, a.candidate_id, a.user_id, a.stage, a.notes, a.draft_id, a.drafted_at, a.applied_date, a.updated_at,
                    a.apply_method, a.client_company, a.recruiter_name, a.recruiter_email AS app_recruiter_email, a.recruiter_phone, a.rate,
                    a.resume_version_id, a.follow_up_date, a.interview_at, a.interview_round, a.last_reply_at,
                    j.title AS job_title, j.company AS job_company, j.location AS job_location, j.url AS job_url, j.source AS job_source,
                    j.recruiter_email AS job_recruiter_email, j.apply_status,
                    c.name AS candidate_name, r.filename AS resume_filename
             FROM applications a JOIN jobs j ON j.id = a.job_id JOIN candidates c ON c.id = a.candidate_id
             LEFT JOIN optimized_resumes r ON r.id = a.resume_version_id"""


def _shape(r) -> Dict:
    d = dict(r)
    for k in ("drafted_at", "applied_date", "updated_at", "last_reply_at"):
        d[k] = str(d[k])[:19] if d.get(k) else ""
    d["recruiter_email"] = d.pop("app_recruiter_email") or d.get("job_recruiter_email") or ""
    d["status_label"] = STATUS_LABELS.get(d["stage"], d["stage"] or "")
    d["apply_method"] = d.get("apply_method") or (d.get("apply_status") or "")
    d["method_label"] = METHODS.get(d["apply_method"], "")
    d["follow_up_due"] = bool(d.get("follow_up_date") and d["follow_up_date"] <= date.today().isoformat() and d["stage"] in OPEN)
    return d


def get(conn, app_id: int) -> Dict:
    cur = conn.cursor()
    cur.execute(_SELECT + " WHERE a.id = ?", (app_id,))
    r = cur.fetchone()
    if not r:
        raise LookupError("Application not found")
    return _shape(r)


def list_for(conn, candidate_ids: List[int], status: str = "", due_only: bool = False) -> List[Dict]:
    if not candidate_ids:
        return []
    cur = conn.cursor()
    sql = _SELECT + f" WHERE a.candidate_id IN ({', '.join(['?'] * len(candidate_ids))})"
    params = list(candidate_ids)
    if status == "open":
        sql += f" AND a.stage IN ({', '.join(['?'] * len(OPEN))})"
        params += sorted(OPEN)
    elif status:
        sql += " AND a.stage = ?"
        params.append(status)
    cur.execute(sql + " ORDER BY a.updated_at DESC, a.id DESC", params)
    rows = [_shape(r) for r in cur.fetchall()]
    return [r for r in rows if r["follow_up_due"]] if due_only else rows


def summary(rows: List[Dict]) -> Dict:
    week_ago = (datetime.utcnow() - timedelta(days=7)).strftime("%Y-%m-%d")
    applied = [r for r in rows if r["stage"] not in ("Saved", "Drafted")]
    return {
        "total": len(rows),
        "applied_this_week": sum(1 for r in applied if (r["applied_date"] or "") >= week_ago),
        "responded": sum(1 for r in rows if ORDER.get(r["stage"], 0) >= ORDER["Responded"] or r["stage"] == "Screening"),
        "interviews": sum(1 for r in rows if r["stage"] in ("Interviewing", "Offer", "Hired")),
        "placed": sum(1 for r in rows if r["stage"] == "Hired"),
        "follow_ups_due": sum(1 for r in rows if r["follow_up_due"]),
    }


def events(conn, app_id: int) -> List[Dict]:
    cur = conn.cursor()
    cur.execute("""SELECT e.kind, e.detail, e.created_at, u.name AS who FROM application_events e
                   LEFT JOIN users u ON u.id = e.user_id WHERE e.application_id = ? ORDER BY e.id DESC""", (app_id,))
    return [{**dict(r), "created_at": str(r["created_at"])[:16], "who": r["who"] or "System"} for r in cur.fetchall()]


# ---------------------------------------------------------------- Gmail replies (read-only)

def _gmail_date(d: str) -> str:
    try:
        return datetime.strptime(d[:10], "%Y-%m-%d").strftime("%d-%b-%Y")
    except ValueError:
        return (datetime.utcnow() - timedelta(days=60)).strftime("%d-%b-%Y")


def check_replies(conn, cand: Dict, user_id: int, imap_factory=None) -> Dict:
    """Search the consultant's Gmail INBOX (read-only) for mail FROM each open application's recruiter
    address since it was drafted / applied. A hit moves 'Draft created' / 'Applied' to 'Recruiter
    responded' and records when. Nothing in Gmail is changed (EXAMINE = read-only, no message bodies)."""
    account, password = (cand.get("gmail_account") or "").strip(), (cand.get("gmail_app_password") or "").replace(" ", "")
    if not (account and password):
        raise TrackerError("This consultant has no Gmail connected (App Password) - connect it on the Consultants tab.")
    rows = [r for r in list_for(conn, [cand["id"]]) if r["stage"] in ("Drafted", "Applied") and "@" in (r["recruiter_email"] or "")]
    if not rows:
        return {"checked": 0, "replied": []}
    try:
        mail = (imap_factory or imaplib.IMAP4_SSL)("imap.gmail.com", 993)   # looked up at call time
        mail.login(account, password)
        mail.select("INBOX", readonly=True)
    except Exception as ex:
        raise TrackerError(f"Could not open {account} (read-only): {str(ex)[:120]}")
    replied = []
    cur = conn.cursor()
    try:
        for r in rows:
            since = _gmail_date(r["applied_date"] or r["drafted_at"] or "")
            addr = r["recruiter_email"].strip().replace('"', "")
            typ, data = mail.search(None, "FROM", f'"{addr}"', "SINCE", since)
            if typ == "OK" and data and data[0].split():
                cur.execute("UPDATE applications SET stage = 'Responded', last_reply_at = CURRENT_TIMESTAMP, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                            (r["id"],))
                event(cur, r["id"], user_id, "reply", f"Reply found in {account} from {addr}")
                replied.append({"id": r["id"], "job_title": r["job_title"], "from": addr})
        conn.commit()
    finally:
        try:
            mail.logout()
        except Exception:
            pass
    return {"checked": len(rows), "replied": replied}


# ---------------------------------------------------------------- Excel

EXPORT_COLUMNS = [("Consultant", "candidate_name"), ("Job title", "job_title"), ("Company", "job_company"), ("Client", "client_company"),
                  ("Status", "status_label"), ("How applied", "method_label"), ("Applied on", "applied_date"),
                  ("Drafted on", "drafted_at"), ("Recruiter", "recruiter_name"), ("Recruiter email", "recruiter_email"),
                  ("Recruiter phone", "recruiter_phone"), ("Rate", "rate"), ("Resume sent", "resume_filename"),
                  ("Follow-up", "follow_up_date"), ("Interview", "interview_at"), ("Round", "interview_round"),
                  ("Last reply", "last_reply_at"), ("Job link", "job_url"), ("Notes", "notes")]


def build_xlsx(rows: List[Dict], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([c for c, _ in EXPORT_COLUMNS])
    for cell in ws[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for r in rows:
        ws.append([r.get(k) or "" for _, k in EXPORT_COLUMNS])
    for i, (name, _) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = max(12, min(40, len(name) + 8))
    ws.freeze_panes = "A3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
