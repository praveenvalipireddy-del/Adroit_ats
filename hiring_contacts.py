"""Saved hiring-team contacts (table hiring_contacts): every person whose email was revealed with Apollo.

- One row per Apollo person (apollo_id). Saved on the reveal - also when Apollo had no VERIFIED email, because
  Apollo charged the credit anyway: the next look at that person shows the result instead of paying again.
- Shared by the whole team and every consultant: a person revealed for one job shows on every job at their company.
- Only an email Apollo marked "verified" is ever stored as `email` (apollo_api.reveal never returns another).
- `emailed_count` / `last_emailed_at` go up when a Gmail draft is created to that address.
"""
import io
import re
from typing import Dict, List, Optional

import apollo_api


def company_norm(name: str) -> str:
    """'Fixture Health Pvt Ltd' and 'FIXTURE HEALTH' -> 'fixture health' (same company, suffixes dropped)."""
    return " ".join(sorted(apollo_api._tokens(name)))


def _shape(r) -> Dict:
    d = dict(r)
    d["verified"] = bool(d.get("verified"))
    for k in ("revealed_at", "last_emailed_at"):
        d[k] = str(d[k])[:19] if d.get(k) else ""
    return d


def by_apollo_ids(conn, ids: List[str]) -> Dict[str, Dict]:
    ids = [i for i in ids if i]
    if not ids:
        return {}
    cur = conn.cursor()
    cur.execute(f"SELECT * FROM hiring_contacts WHERE apollo_id IN ({', '.join(['?'] * len(ids))})", ids)
    return {r["apollo_id"]: _shape(r) for r in cur.fetchall()}


def for_company(conn, company: str) -> List[Dict]:
    """Everyone saved at this company (any job, any consultant), verified emails first."""
    norm = company_norm(company)
    if not norm:
        return []
    cur = conn.cursor()
    first = norm.split()[0]
    cur.execute("SELECT * FROM hiring_contacts WHERE company_norm LIKE ? ORDER BY verified DESC, revealed_at DESC", (f"%{first}%",))
    return [_shape(r) for r in cur.fetchall() if apollo_api.same_company(r["company"] or "", company)]


def save_reveal(conn, person_id: str, res: Dict, company: str, user: Dict, title: str = "", location: str = "") -> Dict:
    """Store what Apollo's match returned (insert, or refresh the row for this person)."""
    cur = conn.cursor()
    row = {"apollo_id": person_id, "company": res.get("company") or company, "company_norm": company_norm(res.get("company") or company),
           "name": res.get("name") or "", "title": res.get("title") or title, "location": location,
           "email": res.get("email") if res.get("verified") else "", "email_status": res.get("email_status") or "",
           "verified": 1 if res.get("verified") else 0, "linkedin_url": safe_linkedin(res.get("linkedin_url")),
           "revealed_by": user["id"], "revealed_by_name": user.get("name") or ""}
    cur.execute("SELECT id FROM hiring_contacts WHERE apollo_id = ?", (person_id,))
    found = cur.fetchone()
    if found:
        cur.execute(f"UPDATE hiring_contacts SET {', '.join(k + ' = ?' for k in row if k != 'apollo_id')} WHERE id = ?",
                    [v for k, v in row.items() if k != "apollo_id"] + [found[0]])
    else:
        cur.execute(f"INSERT INTO hiring_contacts ({', '.join(row)}) VALUES ({', '.join(['?'] * len(row))})", list(row.values()))
    conn.commit()
    return by_apollo_ids(conn, [person_id]).get(person_id) or {}


def mark_emailed(conn, email: str) -> int:
    """A Gmail draft went to this address: count it on the saved contact (if it is one)."""
    email = (email or "").strip().lower()
    if not email:
        return 0
    cur = conn.cursor()
    cur.execute("UPDATE hiring_contacts SET emailed_count = COALESCE(emailed_count, 0) + 1, last_emailed_at = CURRENT_TIMESTAMP "
                "WHERE LOWER(email) = ?", (email,))
    conn.commit()
    return cur.rowcount or 0


def search(conn, q: str = "", company: str = "", limit: int = 500) -> List[Dict]:
    cur = conn.cursor()
    sql, params = "SELECT * FROM hiring_contacts WHERE 1 = 1", []
    if q.strip():
        like = f"%{q.strip().lower()}%"
        sql += " AND (LOWER(name) LIKE ? OR LOWER(company) LIKE ? OR LOWER(email) LIKE ? OR LOWER(title) LIKE ?)"
        params += [like] * 4
    if company.strip():
        sql += " AND company_norm = ?"
        params.append(company_norm(company))
    cur.execute(sql + " ORDER BY revealed_at DESC, id DESC LIMIT ?", params + [limit])
    return [_shape(r) for r in cur.fetchall()]


def companies(conn) -> List[str]:
    cur = conn.cursor()
    cur.execute("SELECT company FROM hiring_contacts GROUP BY company ORDER BY company")
    return [r[0] for r in cur.fetchall() if r[0]]


EXPORT_COLUMNS = [("Name", "name"), ("Title", "title"), ("Company", "company"), ("Location", "location"), ("Email (verified)", "email"),
                  ("LinkedIn", "linkedin_url"), ("Revealed by", "revealed_by_name"), ("Revealed at", "revealed_at"),
                  ("Emailed (drafts)", "emailed_count"), ("Last emailed", "last_emailed_at")]


def build_xlsx(rows: List[Dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    ws.append([c for c, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for r in rows:
        ws.append([r.get(k) if r.get(k) not in (None, "") else "" for _, k in EXPORT_COLUMNS])
    for i, (name, _) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, min(40, len(name) + 10))
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_LINKEDIN_RE = re.compile(r"^https?://([a-z]{2,3}\.)?linkedin\.com/", re.I)


def safe_linkedin(url: Optional[str]) -> str:
    return url if url and _LINKEDIN_RE.match(url) else ""
