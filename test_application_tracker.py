"""Application tracker (application_tracker.py + /api/applications*).

Drafts create a tracked application (real user, recruiter email, resume version sent) and never
move one backwards; the Jobs "Applied" dropdown tracks the row's consultant; manual add for
applications made outside the app; edits are validated and every change is in the history;
summary counts and follow-ups due; Excel export; read-only Gmail reply check (stubbed IMAP); only
the consultant's recruiter or an admin can see or change them; the Reporting board still works.
Gmail is never contacted. Throwaway SQLite DB. Labelled test fixtures.
Run: python test_application_tracker.py
"""
import email
import email.policy
import io
import os
import sys
import tempfile
from datetime import date, timedelta

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_tracker.db")
os.environ["DATABASE_URL"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import application_tracker as at  # noqa: E402
import docx  # noqa: E402
import gmail_multi_manager as gm  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures, captured = [], []


def check(cond, msg):
    if not cond:
        failures.append(msg)


class FakeIMAP:
    """Gmail stub: drafts are appended; the INBOX holds one reply from rec-a@vendor.example."""
    selected = []
    searches = []

    def __init__(self, host, port=None):
        pass

    def login(self, user, pw):
        pass

    def append(self, folder, flags, when, raw):
        captured.append(email.message_from_bytes(raw, policy=email.policy.default))

    def select(self, box, readonly=False):
        FakeIMAP.selected.append((box, readonly))
        return "OK", [b"1"]

    def search(self, charset, *criteria):
        FakeIMAP.searches.append(criteria)
        return ("OK", [b"7"]) if '"rec-a@vendor.example"' in criteria else ("OK", [b""])

    def logout(self):
        pass


gm.imaplib.IMAP4_SSL = FakeIMAP
at.imaplib.IMAP4_SSL = FakeIMAP   # same module - never a real Gmail login
assert at.imaplib.IMAP4_SSL is FakeIMAP
uid = lambda u: u["id"] if isinstance(u, dict) else u  # noqa: E731
rec = {"id": uid(models.create_user("Track Rec", "trk-rec@example.invalid", "Trk-1", role="Recruiter")), "name": "Track Rec", "role": "Recruiter"}
other = {"id": uid(models.create_user("Track Other", "trk-other@example.invalid", "Trk-2", role="Recruiter")), "name": "Track Other", "role": "Recruiter"}
admin = {"id": uid(models.create_user("Track Admin", "trk-admin@example.invalid", "Trk-3", role="Admin")), "name": "Track Admin", "role": "Admin"}
d = docx.Document()
d.add_paragraph("VISAKH TRACKTEST resume fixture")
buf = io.BytesIO()
d.save(buf)
cid = models.create_candidate("Visakh Tracktest", "visakh-trk@example.invalid", title="DevOps Engineer", gmail_account="visakh-trk@example.invalid",
                              assigned_user_id=rec["id"])
models.update_candidate(cid, gmail_app_password="test-app-password-not-real", country="United States")
models.save_resume_file(cid, "Visakh_original.docx", buf.getvalue())


def job(title, rec_email):
    j = models.save_or_update_scraped_job({"title": title, "company": "Vendor Test Co", "source": "Dice", "recruiter_email": rec_email,
                                          "url": f"https://www.dice.com/job-detail/trk-{abs(hash(title))}", "description": "DevOps"})
    return j["id"] if isinstance(j, dict) else j


ja, jb, jc = job("DevOps A Tracktest", "rec-a@vendor.example"), job("DevOps B Tracktest", "rec-b@vendor.example"), job("DevOps C Tracktest", "")
client = app_module.app.test_client()


def login(u):
    with client.session_transaction() as s:
        s["user"] = u


def apps(**q):
    return client.get("/api/applications", query_string=q).get_json()


# ---- a draft creates a tracked application (real user, recruiter email, history)
login(rec)
r = client.post("/api/outreach/create-draft", json={"candidate_id": cid, "job_id": ja}).get_json()
check(r.get("success"), f"draft: {r}")
lst = apps(candidate_id=cid)
a = next(x for x in lst["applications"] if x["job_id"] == ja)
check(a["stage"] == "Drafted" and a["status_label"] == "Draft created" and a["recruiter_email"] == "rec-a@vendor.example" and a["user_id"] == rec["id"],
      f"tracked from the draft: {a}")
ev = client.get(f"/api/applications/{a['id']}/events").get_json()["events"]
check(any(e["kind"] == "created" for e in ev) and any("Created Gmail Draft" in (e["detail"] or "") for e in ev), f"history: {ev}")

# ---- edits: status, fields, validation, history
upd = client.put(f"/api/applications/{a['id']}", json={"status": "Interviewing", "interview_at": "2026-10-20", "interview_round": "Round 1",
                                                       "follow_up_date": (date.today() - timedelta(days=1)).isoformat(), "notes": "Client: Test Bank"}).get_json()
check(upd["application"]["status_label"] == "Interview" and upd["application"]["follow_up_due"], f"updated: {upd}")
check(client.put(f"/api/applications/{a['id']}", json={"status": "Bogus"}).status_code == 400, "unknown status refused")
check(client.put(f"/api/applications/{a['id']}", json={"follow_up_date": "next week"}).status_code == 400, "bad date refused")
ev = client.get(f"/api/applications/{a['id']}/events").get_json()["events"]
check(any("Draft created -> Interview" in e["detail"] and e["who"] == "Track Rec" for e in ev), "status change in history with who")

# ---- a new draft never moves it backwards
client.post("/api/outreach/create-draft", json={"candidate_id": cid, "job_id": ja})
check(next(x for x in apps(candidate_id=cid)["applications"] if x["job_id"] == ja)["stage"] == "Interviewing", "re-draft keeps Interview")

# ---- Jobs "Applied" dropdown tracks the row's consultant
r = client.post(f"/api/jobs/{jb}/apply-status", json={"status": "portal", "candidate_id": cid}).get_json()
b = next(x for x in apps(candidate_id=cid)["applications"] if x["job_id"] == jb)
check(r["success"] and b["stage"] == "Applied" and b["method_label"] == "Applied in portal" and b["applied_date"], f"Applied dropdown -> tracked: {b}")

# ---- manual add
r = client.post("/api/applications", json={"candidate_id": cid, "job_title": "Platform Engineer Tracktest", "company": "Outside Corp",
                                           "job_url": "https://careers.example.invalid/123", "recruiter_email": "hr@outside.example",
                                           "status": "Applied", "apply_method": "both"})
check(r.status_code == 200 and r.get_json()["application"]["method_label"] == "Portal + email", f"manual add: {r.get_json()}")
check(client.post("/api/applications", json={"candidate_id": cid, "job_title": "", "company": "X"}).status_code == 400, "title required")
check(client.post("/api/applications", json={"candidate_id": cid, "job_title": "X", "company": "Y", "job_url": "javascript:alert(1)"}).status_code == 400,
      "only http(s) links")

# ---- summary + filters
lst = apps(candidate_id=cid)
s = lst["summary"]
check(s["total"] == 3 and s["applied_this_week"] == 3 and s["interviews"] == 1 and s["follow_ups_due"] == 1, f"summary: {s}")
check([x["job_id"] for x in apps(candidate_id=cid, due="1")["applications"]] == [ja], "follow-ups due filter")
check(len(apps(candidate_id=cid, status="Applied")["applications"]) == 2, "status filter")

# ---- Gmail reply check (read-only): rec-a replied - but ja is at Interview (not Applied/Drafted), so: add a fresh one
jd_ = job("DevOps D Tracktest", "rec-a@vendor.example")
client.post("/api/outreach/create-draft", json={"candidate_id": cid, "job_id": jd_})
r = client.post("/api/applications/check-replies", json={"candidate_id": cid}).get_json()
check(r.get("success") and [x["from"] for x in r.get("replied", [])] == ["rec-a@vendor.example"] and r["checked"] == 3, f"replies: {r}")
check(all(ro for _, ro in FakeIMAP.selected) and FakeIMAP.selected, f"INBOX opened read-only: {FakeIMAP.selected}")
dapp = next(x for x in apps(candidate_id=cid)["applications"] if x["job_id"] == jd_)
check(dapp["status_label"] == "Recruiter responded" and dapp["last_reply_at"], f"marked responded: {dapp}")
nog = models.create_candidate("NoGmail Tracktest", "nog-trk@example.invalid", assigned_user_id=rec["id"])
check(client.post("/api/applications/check-replies", json={"candidate_id": nog}).status_code == 400, "no Gmail -> clear error")

# ---- Excel
x = client.get(f"/api/applications/export-xlsx?candidate_id={cid}")
ws = load_workbook(io.BytesIO(x.data)).active
check(x.status_code == 200 and ws["A2"].value == "Consultant" and ws.max_row == 2 + 4, f"excel rows: {ws.max_row}")
check("Applications_Visakh_Tracktest.xlsx" in x.headers.get("Content-Disposition", ""), "excel name")

# ---- access
login(other)
check(apps(candidate_id=cid).get("error") == "Consultant not found", "other recruiter can't list")
check(client.put(f"/api/applications/{a['id']}", json={"notes": "x"}).status_code == 404, "other recruiter can't edit")
check(apps()["applications"] == [], "other recruiter's 'all' list doesn't include Visakh")
login(admin)
check(len(apps()["applications"]) == 4, "admin sees all")
with client.session_transaction() as s_:
    s_.clear()
check(client.get("/api/applications").status_code == 401, "login required")

# ---- the Reporting board still reads the same rows
login(rec)
pipe = client.get("/api/pipeline").get_json()
check(isinstance(pipe, dict) and any(len(v) for v in pipe.values()), f"pipeline still works: {list(pipe)[:6]}")
check(at.STATUS_LABELS["Hired"] == "Placed", "board key Hired shows as Placed")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: application tracker - drafts tracked (forward only), Applied dropdown, manual add, edits + history, summary/filters, read-only Gmail replies, Excel, access, board intact.")
