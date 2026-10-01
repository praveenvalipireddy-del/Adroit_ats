"""Filter A / Filter B, ingestion, Excel export and the admin 'unmapped institutions' flow.
Standalone script on a throwaway SQLite DB. No network, no paid calls. All people below are
clearly-labelled test fixtures with example.invalid URLs - not real candidates.
Run: python test_education_filters.py
"""
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_filters.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
config.APIFY_API_TOKEN = ""
import app as app_module  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def edu(school, degree, field="Computer Science", end=None, start=None):
    e = {"schoolName": school, "degree": degree, "fieldOfStudy": field}
    if end:
        e["endDate"] = {"year": end}
    if start:
        e["startDate"] = {"year": start}
    return e


def item(slug, first, education, headline="Software Engineer", company="Test Co", city="Dallas, Texas"):
    return {"linkedinUrl": f"https://www.linkedin.com/in/{slug}.example.invalid/", "firstName": first, "lastName": "Testcase",
            "headline": headline, "location": {"linkedinText": city, "countryCode": "US"},
            "currentPosition": [{"position": headline, "companyName": company}], "education": education}


ADMIN_ID = models.create_user("Filter Test Admin", "filter-admin@example.invalid", "Test-Only-Pass-1", role="Admin")
REC_ID = models.create_user("Filter Test Recruiter", "filter-rec@example.invalid", "Test-Only-Pass-2", role="Recruiter")
if not isinstance(ADMIN_ID, int):
    ADMIN_ID = models.authenticate_user("filter-admin@example.invalid", "Test-Only-Pass-1")["id"]
if not isinstance(REC_ID, int):
    REC_ID = models.authenticate_user("filter-rec@example.invalid", "Test-Only-Pass-2")["id"]

items = [
    item("t-p1", "One", [edu("JNTUH College of Engineering Hyderabad", "Bachelor of Technology - BTech", end=2019),
                         edu("The University of Texas at Dallas", "Master of Science - MS", end=2021)]),
    # US-bound? No: Master's from an unmapped foreign school -> unknown country -> excluded.
    item("t-p2", "Two", [edu("Jawaharlal Nehru Technological University Hyderabad", "B.Tech", end=2016),
                         edu("University of Leeds", "MSc", end=2018)]),
    # multi-degree: Indian Master's (doesn't count) + US Master's (counts)
    item("t-p3", "Three", [edu("Osmania University", "BE", end=2015), edu("IIT Madras", "M.Tech", end=2017),
                           edu("Northeastern University", "Master of Science - MS", end=2020)],
         headline="Data Engineer", company="Acme Test Analytics", city="Boston, Massachusetts"),
    # Bachelor's at an unknown college -> excluded from Filter B until an admin maps it
    item("t-p4", "Four", [edu("Sri Test Engineering College Nowhere", "B.Tech", end=2018),
                          edu("UT Dallas", "MS", end=2020)]),
    # Bachelor's with no end year -> counts only when no year range is set
    item("t-p5", "Five", [edu("JNTUH", "BTech"), edu("University of North Texas", "MS", end=2022)]),
    # dual degree -> "Other", never a Bachelor's
    item("t-p6", "Six", [edu("JNTUH", "Integrated B.Tech + M.Tech (Dual Degree)", end=2019),
                         edu("UT Dallas", "MS", end=2021)]),
]
stats = li.ingest_profiles(li.ApifyProfileProvider().to_profiles(items), "apify", REC_ID)
check(stats == {"added": 6, "updated": 0, "skipped": 0}, f"ingest stats {stats}")

# same person again, different URL spelling -> deduplicated, not a new row
again = li.ingest_profiles(li.ApifyProfileProvider().to_profiles(
    [dict(items[0], linkedinUrl="HTTPS://www.LinkedIn.com/in/T-P1.example.invalid?trk=x")]), "apify", REC_ID)
check(again["added"] == 0 and again["updated"] == 1, f"dedup on LinkedIn URL failed: {again}")

client = app_module.app.test_client()


def login(user_id, name, role):
    with client.session_transaction() as s:
        s["user"] = {"id": user_id, "name": name, "role": role, "email": "x@example.invalid"}


login(REC_ID, "Filter Test Recruiter", "Recruiter")


def names(filter_, inst, **kw):
    q = {"filter": filter_, "institution_id": inst, **kw}
    r = client.get("/api/education/search", query_string=q)
    check(r.status_code == 200, f"search {q} -> {r.status_code} {r.get_json()}")
    return sorted(x["name"] for x in (r.get_json() or {}).get("results", []))


# ---- Filter A: Indian college -> US Master's
check(names("A", "in-jntu-hyderabad") == ["Five Testcase", "One Testcase"],
      f"A/JNTUH: {names('A', 'in-jntu-hyderabad')} (p2 unknown-country MS, p6 dual degree must be excluded)")
check(names("A", "in-jntu-hyderabad", year_from=2018, year_to=2020) == ["One Testcase"],
      "A/JNTUH 2018-2020: missing year must not count as in range")
check(names("A", "in-jntu-hyderabad", year_from=2020) == [], "A/JNTUH from 2020 should be empty")
check(names("A", "in-osmania") == ["Three Testcase"], "A/Osmania multi-degree")
check(names("A", "in-osmania", keyword="data engineer") == ["Three Testcase"], "keyword on headline")
check(names("A", "in-osmania", keyword="acme test") == ["Three Testcase"], "keyword on company")
check(names("A", "in-osmania", keyword="nurse") == [], "keyword mismatch")
check(names("A", "in-osmania", location="boston") == ["Three Testcase"], "location filter")
check(names("A", "in-osmania", location="dallas") == [], "location mismatch")

# ---- Filter B: US university -> Indian undergrad
check(names("B", "us-ut-dallas") == ["One Testcase"], f"B/UTD: {names('B', 'us-ut-dallas')} (p4 unknown college, p6 dual)")
check(names("B", "us-northeastern") == ["Three Testcase"], "B/Northeastern")
check(names("B", "us-ut-dallas", year_from=2020, year_to=2020) == [], "B/UTD 2020 (p1's MS is 2021)")

# ---- result columns
r = client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-osmania"}).get_json()["results"][0]
check("Osmania University (2015)" == r["indian_college"], f"indian_college column: {r['indian_college']}")
check("Northeastern University - Master of Science - MS, Computer Science (2020)" == r["us_masters"],
      f"us_masters column (Indian M.Tech must not appear): {r['us_masters']}")
check(r["linkedin_url"] == "https://www.linkedin.com/in/t-p3.example.invalid", f"url: {r['linkedin_url']}")
check(r["captured_by"] == "Filter Test Recruiter" and r["current_company"] == "Acme Test Analytics", f"captured/company: {r}")

# ---- pagination
p1 = client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-jntu-hyderabad", "page_size": 1}).get_json()
p2 = client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-jntu-hyderabad", "page_size": 1, "page": 2}).get_json()
check(p1["total"] == 2 and p1["pages"] == 2 and len(p1["results"]) == 1 and len(p2["results"]) == 1
      and p1["results"][0]["name"] != p2["results"][0]["name"], f"pagination: {p1} / {p2}")

# ---- validation
check(client.get("/api/education/search", query_string={"filter": "C", "institution_id": "x"}).status_code == 400, "bad filter")
check(client.get("/api/education/search", query_string={"filter": "A"}).status_code == 400, "missing institution")
check(client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-osmania",
                                                         "year_from": 2021, "year_to": 2019}).status_code == 400, "reversed years")

# ---- autocomplete lists
insts = client.get("/api/education/institutions", query_string={"filter": "A"}).get_json()["institutions"]
jntuh = next((i for i in insts if i["id"] == "in-jntu-hyderabad"), None)
check(jntuh and "JNTUH" in jntuh["aliases"] and jntuh["profiles"] == 2, f"A list JNTUH entry: {jntuh}")
check(all(i["id"].startswith("in-") for i in insts), "Filter A list must only hold Indian institutions")
insts_b = client.get("/api/education/institutions", query_string={"filter": "B"}).get_json()["institutions"]
check(insts_b and all(i["id"].startswith("us-") for i in insts_b), "Filter B list must only hold US institutions")

# ---- Excel export (logged with the recruiter id)
r = client.get("/api/education/export-xlsx", query_string={"filter": "A", "institution_id": "in-jntu-hyderabad"})
check(r.status_code == 200 and "spreadsheetml" in r.headers.get("Content-Type", ""), f"xlsx: {r.status_code} {r.headers.get('Content-Type')}")
wb = load_workbook(io.BytesIO(r.data))
ws = wb["Candidates"]
header = [c.value for c in ws[1]]
check(header[:4] == ["Name", "Headline / Current Title", "Current Company", "Location"] and "LinkedIn URL" in header, f"xlsx header {header}")
check(ws.max_row == 3, f"xlsx should have 2 data rows, has {ws.max_row - 1}")
url_col = header.index("LinkedIn URL") + 1
check(ws.cell(row=2, column=url_col).hyperlink is not None, "LinkedIn URL cell should be a clickable link")
conn = models.get_db_connection()
cur = conn.cursor()
cur.execute("SELECT user_id, details FROM capture_log WHERE action = 'export_xlsx'")
logs = [tuple(x) for x in cur.fetchall()]
check(len(logs) == 1 and logs[0][0] == REC_ID and logs[0][1].startswith("2 rows"), f"export log: {logs}")
cur.execute("SELECT COUNT(*) FROM capture_log WHERE action = 'capture' AND user_id = ?", (REC_ID,))
check(cur.fetchone()[0] == 6, "each captured profile should be logged with the recruiter id")
conn.close()

# ---- admin: unmapped list + mapping
check(client.get("/api/education/unmapped").status_code == 403, "recruiter must not see the admin list")
login(ADMIN_ID, "Filter Test Admin", "Admin")
data = client.get("/api/education/unmapped").get_json()
unm = {i["name"]: i for i in data["items"]}
check("Sri Test Engineering College Nowhere" in unm and "University of Leeds" in unm, f"unmapped list: {list(unm)}")
check(unm.get("Sri Test Engineering College Nowhere", {}).get("profiles") == 1, "unmapped list should count profiles")
r = client.post(f"/api/education/unmapped/{unm['Sri Test Engineering College Nowhere']['id']}/map",
                json={"new_name": "Sri Test Engineering College", "country": "India", "city": "Testville"})
check(r.status_code == 200 and r.get_json().get("entries_updated") == 1, f"map new institution: {r.status_code} {r.get_json()}")
check(names("B", "us-ut-dallas") == ["Four Testcase", "One Testcase"], "after mapping, p4 must appear in Filter B")
r = client.post(f"/api/education/unmapped/{unm['University of Leeds']['id']}/map", json={"canonical_id": "us-ut-dallas"})
check(r.status_code == 200, f"map to existing: {r.status_code} {r.get_json()}")
check(names("A", "in-jntu-hyderabad") == ["Five Testcase", "One Testcase", "Two Testcase"], "after mapping Leeds (to test UTD), p2 counts")
bad = client.post("/api/education/unmapped/999999/map", json={"canonical_id": "us-ut-dallas"})
check(bad.status_code == 404, "mapping a missing item")

# a NEW profile with the mapped name is matched straight away (cache refreshed)
li.ingest_profiles(li.ApifyProfileProvider().to_profiles([item("t-p7", "Seven", [
    edu("Sri Test Engineering College Nowhere", "B.Tech", end=2017), edu("UT Dallas", "MS", end=2019)])]), "apify", REC_ID)
check("Seven Testcase" in names("B", "us-ut-dallas"), "new profile with an admin-mapped name should match")

# ---- sourcing-pool import (startup) + upgrade by a full profile
conn = models.get_db_connection()
cur = conn.cursor()
cur.execute("""INSERT INTO sourced_candidates (source, search_year, profile_url, name, headline, bachelor_year, bachelor_degree,
               bachelor_college, master_degree, master_university, master_year, location)
               VALUES ('pdl', '2017', 'https://www.linkedin.com/in/t-pool.example.invalid', 'Pool Testcase', 'QA Engineer', '2017',
                       'bachelors - computer science', 'Osmania University, Hyderabad', 'masters - computer science (2019)',
                       'University of North Texas', '2019', 'Denton, Texas')""")
conn.commit()
conn.close()
models.init_db()
models.init_db()
check(names("A", "in-osmania") == ["Pool Testcase", "Three Testcase"], f"pool import: {names('A', 'in-osmania')}")
conn = models.get_db_connection()
cur = conn.cursor()
cur.execute("SELECT COUNT(*), MAX(education_complete) FROM linkedin_profiles WHERE linkedin_url LIKE '%t-pool%'")
cnt, complete = cur.fetchone()
check(cnt == 1 and complete == 0, f"pool profile imported once, marked incomplete: {cnt}, {complete}")
conn.close()
li.ingest_profiles(li.ApifyProfileProvider().to_profiles([item("t-pool", "Pool", [
    edu("Osmania University", "BE", end=2017), edu("University of North Texas", "MS", end=2019),
    edu("Northeastern University", "MBA", end=2023)])]), "apify", REC_ID)
check("Pool Testcase" in names("B", "us-northeastern"), "full profile should replace the pool-only education")
# ...and a later pool row for the same person must NOT downgrade it back
li.ingest_profiles(li.SourcedPoolProvider().to_profiles([{"profile_url": "https://www.linkedin.com/in/t-pool.example.invalid",
                   "name": "Pool Testcase", "bachelor_college": "Osmania University", "bachelor_year": "2017"}]), "apify", None)
check("Pool Testcase" in names("B", "us-northeastern"), "pool data must not overwrite a complete profile")

# ---- a live Apify search poll stores EVERY scanned profile for the filters (Apify call replaced by a stub)
import linkedin_sourcing  # noqa: E402

scanned = [item("t-live", "Live", [edu("VIT University", "B.Tech", end=2018), edu("NJIT", "MS", end=2020)])]


def fake_poll(run_id, dataset_id, bachelor_year, offset, matched_so_far, target=None):
    return {"status": "SUCCEEDED", "done": True, "next_offset": 1, "scanned_new": 1, "new_matches": [],
            "other_year_matches": [], "skipped": {"wrong_bachelor_year": 1}, "raw_items": scanned}


real_poll = linkedin_sourcing.poll_search
linkedin_sourcing.poll_search = fake_poll
try:
    login(REC_ID, "Filter Test Recruiter", "Recruiter")
    r = client.post("/api/students/search-poll", json={"run_id": "AbCdEfGhIjKlMnOpQ", "dataset_id": "QpOnMlKjIhGfEdCbA", "bachelor_year": 2023})
finally:
    linkedin_sourcing.poll_search = real_poll
check(r.status_code == 200, f"search-poll: {r.status_code} {r.get_json()}")
check("raw_items" not in (r.get_json() or {}), "raw scanned profiles must not be sent to the browser")
check(names("A", "in-vit") == ["Live Testcase"], "a profile scanned by a live search (not a match for its year) should be in Filter A")
check(names("B", "us-njit") == ["Live Testcase"], "...and in Filter B")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: Filter A/B (multi-degree, unknown country, years, location/keyword, paging), dedup, Excel export + log, admin mapping, pool import/upgrade.")
