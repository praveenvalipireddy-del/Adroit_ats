"""Browser test (Playwright + Microsoft Edge) for the Education Filters panel and the admin
"Colleges not recognised" list.

Real Flask app on a throwaway SQLite DB, real login form. Checks: autocomplete by alias (JNTUH),
Filter A results + columns, paging (27 profiles -> 2 pages), year range, Filter B tab, the actual
Excel download (opened with openpyxl), and admin linking an unrecognised college so its profile
appears in Filter B. No JavaScript errors allowed.

COST SAFETY: every request to api.apify.com and the paid search endpoints is aborted; the Apify
token is a dummy. All people are labelled test fixtures with example.invalid URLs.
Run: python test_education_filters_browser.py
"""
import os
import socket
import sys
import tempfile
import threading

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_edu_ui.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = "dummy-test-token-not-real"
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
config.APIFY_API_TOKEN = "dummy-test-token-not-real"
import app as app_module  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from werkzeug.serving import make_server  # noqa: E402

EMAIL, PASSWORD = "edu-ui-admin@example.invalid", "EduUi-Test-Only-1"
admin = models.create_user("Edu UI Admin", EMAIL, PASSWORD, role="Admin")
admin_id = admin["id"] if isinstance(admin, dict) else admin


def item(slug, first, education):
    return {"linkedinUrl": f"https://www.linkedin.com/in/{slug}.example.invalid", "firstName": first, "lastName": "Uitest",
            "headline": "Software Engineer", "location": {"linkedinText": "Dallas, Texas"},
            "currentPosition": [{"position": "Software Engineer", "companyName": "Test Corp"}], "education": education}


def e(school, degree, end):
    return {"schoolName": school, "degree": degree, "endDate": {"year": end}}


items = [item(f"jn-{i:02d}", f"Jntuh{i:02d}", [e("JNTUH", "B.Tech", 2015 + i % 5), e("UT Dallas", "MS", 2021)]) for i in range(27)]
items.append(item("unk-1", "Unknownbachelor", [e("Sri Ui Test College Faraway", "B.Tech", 2016), e("UT Dallas", "MS", 2019)]))
li.ingest_profiles(li.ApifyProfileProvider().to_profiles(items), "apify", admin_id)

with socket.socket() as s:
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
server = make_server("127.0.0.1", port, app_module.app, threaded=True)
threading.Thread(target=server.serve_forever, daemon=True).start()
base = f"http://127.0.0.1:{port}"

failures, js_errors, blocked = [], [], []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def block(route):
    blocked.append(route.request.url)
    route.abort()


shot = os.path.join(tmp_dir, "edu_filters.png")
try:
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="msedge", headless=True)
        ctx = browser.new_context(accept_downloads=True, viewport={"width": 1440, "height": 1000})
        page = ctx.new_page()
        page.on("pageerror", lambda err: js_errors.append(str(err)))
        page.route("**/api.apify.com/**", block)
        page.route("**/api/students/search-start**", block)
        page.route("**/api/students/search-poll**", block)

        page.goto(base + "/login")
        page.fill("input[name=email]", EMAIL)
        page.fill("input[name=password]", PASSWORD)
        page.click("button[type=submit]")
        page.wait_for_url("**/dashboard**")
        check("app.js?v=5.36.0" in page.content(), "cache-buster not bumped to 5.36.0")

        page.click("a.nav-item[data-tab=sourcing]")
        page.wait_for_selector("#edu-panel", state="visible")
        # every control must sit inside the panel (nothing spilling past its right edge) at common widths
        for width in (1280, 1440, 1024):
            page.set_viewport_size({"width": width, "height": 1000})
            overflow = page.evaluate("""() => { const panel = document.getElementById('edu-panel').getBoundingClientRect();
                return [...document.querySelectorAll('#form-edu-search > *')].filter(el => el.getBoundingClientRect().right > panel.right + 1).map(el => el.id || el.tagName); }""")
            check(not overflow, f"at {width}px these controls spill outside the panel: {overflow}")
        page.set_viewport_size({"width": 1440, "height": 1000})
        check(page.is_disabled("#btn-edu-export"), "Excel button should be disabled before a search")

        # searching before choosing a college explains what to do
        page.click("#btn-edu-search")
        check("Pick an Indian college" in page.inner_text("#edu-status"), "no-college message missing")

        # autocomplete by alias
        page.fill("#edu-inst-input", "jntuh")
        page.wait_for_selector(".edu-inst-option")
        first = page.inner_text(".edu-inst-option")
        check("Jawaharlal Nehru Technological University Hyderabad" in first and "27 candidates" in first, f"autocomplete: {first!r}")
        page.click(".edu-inst-option")
        page.wait_for_function("document.querySelectorAll('#edu-results tr').length > 0")
        check("27" in page.inner_text("#edu-status"), f"status: {page.inner_text('#edu-status')}")
        check(page.locator("#edu-results tr").count() == 25, "page 1 should show 25 rows")
        check(page.inner_text("#edu-page-info") == "Page 1 of 2", f"pager: {page.inner_text('#edu-page-info')}")
        row = page.inner_text("#edu-results tr:first-child")
        check("UT Dallas" in row or "University of Texas at Dallas" in row, f"US Master's column missing: {row!r}")
        check("Edu UI Admin" in row and "Test Corp" in row, f"captured by / company column: {row!r}")
        href = page.get_attribute("#edu-results tr:first-child a", "href")
        check(href and href.startswith("https://www.linkedin.com/in/jn-"), f"LinkedIn link: {href}")
        page.screenshot(path=shot)

        page.click("#btn-edu-next")
        page.wait_for_function("document.getElementById('edu-page-info').textContent === 'Page 2 of 2'")
        check(page.locator("#edu-results tr").count() == 2, "page 2 should show 2 rows")

        # year range: 2015 + i%5 == 2017 -> i in {2,7,12,17,22} -> 5 people
        page.fill("#edu-year-from", "2017")
        page.fill("#edu-year-to", "2017")
        page.click("#btn-edu-search")
        page.wait_for_function("document.getElementById('edu-status').innerText.includes('5 candidates')")

        # Excel download of the current (filtered) search
        with page.expect_download() as dl_info:
            page.click("#btn-edu-export")
        path = os.path.join(tmp_dir, "download.xlsx")
        dl_info.value.save_as(path)
        wb = load_workbook(path)
        ws = wb["Candidates"]
        check(ws.max_row == 6, f"Excel should have 5 data rows, has {ws.max_row - 1}")
        check(dl_info.value.suggested_filename.endswith(".xlsx"), f"filename {dl_info.value.suggested_filename}")

        # Filter B tab
        page.click(".edu-tab[data-filter=B]")
        check(page.text_content("#edu-inst-label").startswith("US university"), f"Filter B label: {page.text_content('#edu-inst-label')!r}")
        page.fill("#edu-year-from", "")
        page.fill("#edu-year-to", "")
        page.fill("#edu-inst-input", "UT Dallas")
        page.wait_for_selector(".edu-inst-option")
        page.click(".edu-inst-option")
        page.wait_for_function("document.getElementById('edu-status').innerText.includes('27 candidates')")
        check("Unknownbachelor" not in page.inner_text("#edu-results"), "unrecognised college must not count yet")

        # Admin: link the unrecognised college as a new Indian college
        page.click("a.nav-item[data-tab=settings]")
        page.wait_for_selector("#unmapped-body tr")
        row = page.locator("#unmapped-body tr", has_text="Sri Ui Test College Faraway")
        check(row.count() == 1, "unrecognised college not listed for the admin")
        row.locator(".unmapped-new-country").select_option("India")
        row.locator(".unmapped-new-city").fill("Testpur")
        row.locator(".unmapped-create").click()
        page.wait_for_function("!document.getElementById('unmapped-body').innerText.includes('Sri Ui Test College Faraway')")

        # back to Filter B: the profile now counts
        page.click("a.nav-item[data-tab=sourcing]")
        page.wait_for_selector(".edu-tab[data-filter=A]", state="visible")
        page.click(".edu-tab[data-filter=A]")
        page.click(".edu-tab[data-filter=B]")
        page.wait_for_function("document.getElementById('edu-status').innerText.includes('28 candidates')")
        check("Unknownbachelor" in page.inner_text("#edu-results"), "mapped profile should now appear in Filter B")
        browser.close()
finally:
    server.shutdown()

check(not js_errors, f"JavaScript errors: {js_errors}")
check(not blocked, f"a paid Apify call was attempted (blocked): {blocked}")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    print("screenshot:", shot)
    sys.exit(1)
print(f"PASS: Education Filters in Edge - autocomplete, A/B, paging, years, Excel download, admin linking, no JS errors, no paid calls. Screenshot: {shot}")
