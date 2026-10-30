"""The job page (Option A) and Contacts, in Edge - Apollo stubbed (no network, no credits), labelled fixtures.

1. /api/jobs/<id>/page: the job, the consultant's checklist (resume on file, tailored version for this job, draft,
   applied), the saved hiring team at the company, the company's other jobs; login + consultant access.
2. Edge: "Open job" on the consultant's Matched jobs opens the page: header, match for the consultant, checklist
   2 of 4, the tailored resume chosen by default, Preview renders the Word file.
3. Hiring team tab: the free search lists people, a saved person first with the email and LinkedIn, "Get verified
   email" asks before the credit, a guessed email is never shown; ticking a saved person -> one Gmail draft with the
   chosen resume. Company tab: the company's other jobs. Back returns to the consultant.
4. Contacts: every saved person; Excel export for admins only.
Run: python test_job_page_browser.py
"""
import io
import os
import socket
import sys
import tempfile
import threading

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_job_page.db")
os.environ["DATABASE_URL"] = ""
os.environ["APOLLO_API_KEY"] = "fixture-not-a-real-key"
for k in ("GEMINI_API_KEY", "ANTHROPIC_API_KEY", "OPENROUTER_API_KEY", "XAI_API_KEY", "APIFY_API_TOKEN"):
    os.environ[k] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import docx  # noqa: E402
import apollo_api  # noqa: E402
import app as app_module  # noqa: E402
import hiring_contacts  # noqa: E402
import models  # noqa: E402
import resume_versions  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from werkzeug.serving import make_server  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


P1, P2, P3 = "a" * 24, "b" * 24, "c" * 24
calls = []


def fake_post(path, body):
    calls.append(path)
    if path == "/mixed_people/api_search":
        return {"people": [
            {"id": P2, "first_name": "Hari", "last_name_obfuscated": "H***", "title": "HR Business Partner", "organization": {"name": "Pagefixture Labs"}},
            {"id": P1, "first_name": "Rita", "last_name_obfuscated": "Re***r", "title": "Senior Technical Recruiter", "organization": {"name": "Pagefixture Labs"}},
            {"id": P3, "first_name": "Mona", "last_name_obfuscated": "M***", "title": "Talent Acquisition Lead", "organization": {"name": "Pagefixture Labs"},
             "has_email": False}]}
    if path == "/people/match":
        if body["id"] == P2:
            return {"person": {"name": "Hari Hfixture", "email": "hari@pagefixture.example", "email_status": "guessed", "organization": {"name": "Pagefixture Labs"}}}
        return {"person": None}
    raise AssertionError(path)


apollo_api.http_post = fake_post

uid = lambda u: u["id"] if isinstance(u, dict) else u  # noqa: E731
EMAIL, PW = "jp-rec@example.invalid", "JobPage-1"
REC = uid(models.create_user("Jaya Pagetest", EMAIL, PW, role="Recruiter", region="USA"))
ADMIN_EMAIL, ADMIN_PW = "jp-admin@example.invalid", "JobPage-Admin-1"
ADM = uid(models.create_user("Ada Pageadmin", ADMIN_EMAIL, ADMIN_PW, role="Admin", region="USA"))
OTHER = uid(models.create_user("Otto Otherrec", "jp-other@example.invalid", "JobPage-2", role="Recruiter", region="USA"))
K = models.create_candidate("Kavi Pagefixture", "kavi.page@example.invalid", title="Data Analyst", primary_skills="SQL, Power BI, Python, Tableau",
                            experience_years=5, country="United States", assigned_user_id=REC,
                            resume_text="KAVI PAGEFIXTURE\nData Analyst - SQL, Power BI, Python")
d = docx.Document()
for line in ["KAVI PAGEFIXTURE", "Data Analyst - SQL, Power BI, Python, tailored for Pagefixture (fixture)"]:
    d.add_paragraph(line)
buf = io.BytesIO()
d.save(buf)


def job(title, url):
    j = models.save_or_update_scraped_job({"title": title, "company": "Pagefixture Labs", "source": "Company site (Greenhouse)", "url": url,
                                           "location": "Dallas, TX", "job_type": "Full-time", "country": "United States",
                                           "description": "We need SQL, Power BI and Python for reporting dashboards. Labelled test fixture."})
    return j["id"] if isinstance(j, dict) else j


JID = job("Data Analyst (Job Page Fixture)", "https://job-boards.greenhouse.io/pagefixture/jobs/1")
JID2 = job("BI Developer (Job Page Fixture)", "https://job-boards.greenhouse.io/pagefixture/jobs/2")
VID = resume_versions.save(K, JID, "Kavi_Pagefixture_Data_Analyst.docx", buf.getvalue(), "Data Analyst", "", "fixture", REC)

# one person revealed before (by the admin, for another job) - saved for everyone
conn = models.get_db_connection()
hiring_contacts.save_reveal(conn, P1, {"matched": True, "verified": True, "email": "rita@pagefixture.example", "email_status": "verified",
                                       "name": "Rita Recruiter", "title": "Senior Technical Recruiter", "company": "Pagefixture Labs",
                                       "linkedin_url": "https://www.linkedin.com/in/rita-pagefixture"}, "Pagefixture Labs",
                            {"id": ADM, "name": "Ada Pageadmin"})
conn.close()

# ---- 1. the page route
client = app_module.app.test_client()
check(client.get(f"/api/jobs/{JID}/page").status_code == 401, "login required")
with client.session_transaction() as s_:
    s_["user"] = {"id": REC, "name": "Jaya", "role": "Recruiter", "email": EMAIL}
pg = client.get(f"/api/jobs/{JID}/page?candidate_id={K}").get_json()
check(pg["job"]["title"] == "Data Analyst (Job Page Fixture)" and pg["resume"]["on_file"], f"job + resume: {pg.get('resume')}")
check([v["id"] for v in pg["versions"]] == [VID] and pg["versions"][0]["job_id"] == JID, f"tailored version: {pg['versions']}")
check([c["apollo_id"] for c in pg["contacts"]] == [P1] and pg["contacts"][0]["email"] == "rita@pagefixture.example", f"saved contacts: {pg['contacts']}")
check([j["id"] for j in pg["company_jobs"]] == [JID2], f"the company's other jobs: {pg['company_jobs']}")
check(pg["application"] is None, "no application yet")
with client.session_transaction() as s_:
    s_["user"] = {"id": OTHER, "name": "Otto", "role": "Recruiter", "email": "jp-other@example.invalid"}
pg = client.get(f"/api/jobs/{JID}/page?candidate_id={K}").get_json()
check(pg["resume"] is None and pg["versions"] == [], "another recruiter's consultant: no resume / versions")
check(client.get("/api/contacts/export.xlsx").status_code == 403, "recruiters can't export contacts")
check(client.get("/api/contacts").get_json()["can_export"] is False, "can_export false for a recruiter")
with client.session_transaction() as s_:
    s_["user"] = {"id": ADM, "name": "Ada", "role": "Admin", "email": ADMIN_EMAIL}
x = client.get("/api/contacts/export.xlsx")
check(x.status_code == 200 and x.data[:2] == b"PK", "admin Excel export")
import openpyxl  # noqa: E402
ws = openpyxl.load_workbook(io.BytesIO(x.data)).active
check([c.value for c in ws[1]][:5] == ["Name", "Title", "Company", "Location", "Email (verified)"] and ws.cell(row=2, column=5).value == "rita@pagefixture.example",
      "export columns + the verified email")

# ---- 2-4. Edge
with socket.socket() as sk:
    sk.bind(("127.0.0.1", 0))
    port = sk.getsockname()[1]
server = make_server("127.0.0.1", port, app_module.app, threaded=True)
threading.Thread(target=server.serve_forever, daemon=True).start()
js_errors, shots = [], {}
try:
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="msedge", headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 1000})
        page.on("pageerror", lambda e: js_errors.append(str(e)))
        dialogs = []
        page.on("dialog", lambda dlg: (dialogs.append(dlg.message), dlg.accept()))
        page.goto(f"http://127.0.0.1:{port}/login")
        page.fill("input[name=email]", EMAIL)
        page.fill("input[name=password]", PW)
        page.click("button[type=submit]")
        page.wait_for_url("**/dashboard**")
        check("app.js?v=5.99.14" in page.content(), "cache-buster not bumped to 5.99.14")
        page.evaluate("() => { window.__drafts = []; window.createJobDraft = async (...a) => { window.__drafts.push(a); }; }")
        page.click("a.nav-item[data-tab=candidates]")
        page.click(f".cons-card[data-id='{K}']")
        page.wait_for_selector(f".cons-job[data-id='{JID}'] .cons-job-open", timeout=15000)
        page.click(f".cons-job[data-id='{JID}'] button.btn.cons-job-open")
        page.wait_for_selector("#tab-jobpage.active .jp-step", timeout=10000)
        page.wait_for_function("() => typeof jp !== 'undefined' && jp.data && Array.isArray(jp.data.versions)", timeout=15000)
        txt = page.inner_text("#tab-jobpage")
        check("Data Analyst (Job Page Fixture)" in txt and "Pagefixture Labs" in txt and "MATCH FOR KAVI PAGEFIXTURE" in txt, f"header + match: {txt[:300]!r}")
        check("2 of 4 done" in txt, f"checklist: resume + tailored done: {txt[:900]!r}")
        check(page.eval_on_selector("#jp-resume", "e => e.value") == str(VID), "the tailored resume is chosen by default")
        check("Rita Recruiter" in txt and "rita@pagefixture.example" in txt, "the saved recruiter shows on Overview")
        shots["overview"] = os.path.join(tmp_dir, "jp_overview.png")
        page.screenshot(path=shots["overview"], full_page=True)

        page.click("#jp-resume ~ .jp-row .jp-preview")
        page.wait_for_selector("#jp-preview-body .docx-wrapper", timeout=15000)
        check("tailored for Pagefixture" in page.inner_text("#jp-preview-body"), "Preview renders the tailored Word file")
        page.click("#btn-close-jp-preview")

        # -- Hiring team
        page.click(".jp-tab[data-pane=team]")
        page.wait_for_selector(".jp-trow", timeout=10000)
        rows = page.eval_on_selector_all(".jp-trow", "els => els.map(e => e.dataset.id)")
        check(rows[0] == P1 and set(rows) == {P1, P2, P3}, f"saved person first: {rows}")
        r1 = page.inner_text(f".jp-trow[data-id='{P1}']")
        check("SAVED" in r1 and "rita@pagefixture.example" in r1 and "Ada Pageadmin" in r1, f"saved row: {r1!r}")
        check(page.eval_on_selector(f".jp-trow[data-id='{P1}'] .jp-li", "e => e.href") == "https://www.linkedin.com/in/rita-pagefixture", "LinkedIn profile link")
        check("linkedin.com/search/results/people" in page.eval_on_selector(f".jp-trow[data-id='{P2}'] .jp-li", "e => e.href"), "free LinkedIn search for the others")
        check("No email in Apollo" in page.inner_text(f".jp-trow[data-id='{P3}']"), "no email in Apollo")
        n_match = calls.count("/people/match")
        page.click(f".jp-trow[data-id='{P2}'] .jp-reveal")
        page.wait_for_selector(f".jp-trow[data-id='{P2}'] .sa-warn", timeout=10000)
        check(dialogs and "1 Apollo credit" in dialogs[-1], f"the credit is confirmed first: {dialogs}")
        check(calls.count("/people/match") == n_match + 1 and "hari@" not in page.inner_text("#tab-jobpage"), "guessed email never shown")
        check(page.eval_on_selector(f".jp-trow[data-id='{P2}'] .jp-tick", "e => e.disabled"), "no verified email -> can't be ticked")
        page.check(f".jp-trow[data-id='{P1}'] .jp-tick")
        check("Create 1 Gmail draft" in page.inner_text("#jp-draft-ticked") and "Rita Recruiter" in page.inner_text(".jp-side"), "aside: 1 ticked")
        shots["team"] = os.path.join(tmp_dir, "jp_team.png")
        page.screenshot(path=shots["team"], full_page=True)
        page.click("#jp-draft-ticked")
        page.wait_for_function("window.__drafts.length === 1", timeout=5000)
        check(page.evaluate("window.__drafts[0]") == [JID, K, "rita@pagefixture.example", None, VID], f"draft: job, consultant, email, tailored version: {page.evaluate('window.__drafts')}")

        # -- Company + back
        page.click(".jp-tab[data-pane=company]")
        check("BI Developer (Job Page Fixture)" in page.inner_text("#jp-pane"), "the company's other jobs")
        page.click("#jp-back")
        page.wait_for_selector("#tab-consultants.active", timeout=5000)
        check(True, "back to the consultant")

        # -- Contacts
        page.click("a.nav-item[data-tab=contacts]")
        page.wait_for_selector("#contacts-table .ct-row", timeout=10000)
        ct = page.inner_text("#contacts-table")
        check("Rita Recruiter" in ct and "rita@pagefixture.example" in ct and "Hari Hfixture" in ct and "No verified email" in ct, f"contacts: {ct!r}")
        check(not page.is_visible("#contacts-export"), "no Export button for a recruiter")
        page.fill("#contacts-q", "hari")
        page.wait_for_function("!document.querySelector('#contacts-table').innerText.includes('Rita')", timeout=5000)
        shots["contacts"] = os.path.join(tmp_dir, "jp_contacts.png")
        page.screenshot(path=shots["contacts"], full_page=True)

        # -- a narrow window (tablet): the page stacks, nothing scrolls sideways
        # (at phone width the whole dashboard keeps its 280px menu - not this page's layout)
        page.evaluate("() => openJobPage(state.jobs.find(j => j.id === %d), %d, 'jobs')" % (JID, K))
        page.wait_for_selector("#tab-jobpage.active .jp-step", timeout=10000)
        page.set_viewport_size({"width": 900, "height": 900})
        page.wait_for_timeout(300)
        w = page.evaluate("document.querySelector('#jobpage-root').scrollWidth - document.querySelector('#jobpage-root').clientWidth")
        check(w <= 1, f"no sideways scroll in the job page in a narrow window ({w}px)")
        browser.close()
finally:
    server.shutdown()

check(not js_errors, f"JS errors: {js_errors}")
if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    print("screenshots:", shots)
    sys.exit(1)
print(f"PASS: job page - route + access, checklist, tailored resume chosen + preview, hiring team (saved first, LinkedIn, credit confirmed, "
      f"guessed never shown, ticked -> draft with the tailored version), company jobs, back, Contacts + admin-only export. Screenshots: {shots}")
