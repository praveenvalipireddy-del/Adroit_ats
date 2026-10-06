"""Resume optimizer speed: changes-only rewrite, analysis cache, JD trimming, per-step timings.

Step 2 used to make the AI repeat the WHOLE resume; now it returns only the changed paragraphs
("[12] new text", "[+12] new bullet") and everything else stays exactly as it was. Checks: the
step-2 prompt asks for changes only; the edits land in the original Word file with untouched
paragraphs byte-identical; a "[+N]" bullet is inserted after paragraph N; the same resume + JD
reuses the cached analysis (one AI call instead of two); JD boilerplate (EEO, benefits) is trimmed
but requirements kept; the plain-text path works; RESUME_REWRITE_MODE=full still works; timings are
reported. Also prints how much shorter the AI's step-2 answer is (measured on this fixture).
Gemini is never called: resume_bot._generate is a stub. Labelled test fixture.
Run: python test_optimizer_speed.py
"""
import base64
import io
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.GEMINI_API_KEY = "test-key-not-real"
import docx  # noqa: E402
import docx_editor  # noqa: E402
import resume_bot as rb  # noqa: E402

failures, calls = [], []


def check(cond, msg):
    if not cond:
        failures.append(msg)


# ---- a realistic Word resume: name, summary, skills, 3 jobs with bullets (fixture)
d = docx.Document()
d.add_paragraph("RAVI SPEEDTEST")
d.add_paragraph("Senior Data Engineer | ravi.speedtest@example.invalid | +1 555 0100")
d.add_heading("PROFESSIONAL SUMMARY", level=1)
d.add_paragraph("Data Engineer with 8 years of experience building batch and streaming pipelines on AWS using Spark, Python and SQL.")
d.add_heading("TECHNICAL SKILLS", level=1)
d.add_paragraph("Languages: Python, SQL, Scala")
d.add_paragraph("Big Data: Spark, Hadoop, Hive, Airflow")
d.add_paragraph("Cloud: AWS (S3, EMR, Glue, Lambda, Redshift)")
d.add_heading("PROFESSIONAL EXPERIENCE", level=1)
for company, years in (("Example Bank Testco", "Jan 2022 - Present"), ("Example Retail Testco", "Mar 2019 - Dec 2021"), ("Example Health Testco", "Jun 2017 - Feb 2019")):
    d.add_paragraph(f"{company} | Data Engineer | {years}")
    for i in range(7):
        d.add_paragraph(f"Built and maintained pipeline number {i + 1} at {company} processing daily batch data with Spark and Python on AWS EMR.", style="List Bullet")
    d.add_paragraph("Environment: Spark, Python, SQL, AWS S3, EMR, Glue, Airflow")
d.add_heading("EDUCATION", level=1)
d.add_paragraph("Master of Science in Computer Science, Example State University Testville, 2017")
buf = io.BytesIO()
d.save(buf)
DOCX = buf.getvalue()
_, paras = docx_editor.load_paragraphs(DOCX)
summary_idx = next(p["index"] for p in paras if p["text"].startswith("Data Engineer with 8"))
skills_idx = next(p["index"] for p in paras if p["text"].startswith("Big Data:"))
last_bullet = max(p["index"] for p in paras if "Example Bank Testco processing" in p["text"])

ANALYSIS = {
    "initial_match_percentage": 74,
    "match_breakdown": {"mandatory_skills": 30, "recent_project_relevance": 18, "domain_experience": 12,
                        "tools_frameworks_cloud": 7, "certifications_education": 5, "location_work_authorization": 2},
    "domain_detected": "Banking", "strong_match_skills": ["Spark", "Python"], "partial_match_skills": ["Airflow"],
    "mandatory_missing_skills": ["Kafka"], "preferred_missing_skills": [], "risky_skills_avoided": ["Flink - no background"],
    "skills_to_add": {"summary": ["Kafka"], "technical_skills": ["Kafka"], "projects": ["Kafka"], "environment": ["Kafka"]},
    "ats_optimization_notes": ["Keywords optimized: Kafka"], "target_match_percentage": 86,
}
CHANGES = (f"[{summary_idx}] Data Engineer with 8 years of experience building batch and streaming pipelines on AWS using Spark, Kafka, Python and SQL.\n"
           f"[{skills_idx}] Big Data: Spark, Kafka, Hadoop, Hive, Airflow\n"
           f"[+{last_bullet}] Streamed real-time transaction events with Kafka into Spark jobs for fraud reporting.\n")
prompts = []


def fake_generate(client, types, system, prompt, want_json, max_tokens):
    calls.append("analysis" if want_json else "rewrite")
    prompts.append((system, prompt, max_tokens))
    if want_json:
        return json.dumps(ANALYSIS), False, "stub-model"
    return CHANGES, False, "stub-model"


rb._generate = fake_generate
rb.gemini_configured = lambda: True

JD = """Senior Data Engineer - Banking

Requirements:
- 7+ years of data engineering experience with Spark, Python and Kafka
- AWS experience required

We are an equal opportunity employer. All qualified applicants will receive consideration without regard to race, color, religion.

Benefits include 401(k), dental and vision insurance, paid time off."""

# ---- docx path, changes-only
res = rb.optimize_resume_for_jd("", JD, docx_bytes=DOCX)
check(res.get("optimized") and res.get("format_preserved"), f"optimized in place: {res.get('not_optimized_reason')}")
check(calls == ["analysis", "rewrite"], f"two AI calls: {calls}")
sys_prompt, user_prompt, max_out = prompts[-1]
check("ONLY THE CHANGED PARAGRAPHS" in sys_prompt and "never repeat a paragraph you leave unchanged" in sys_prompt, "step 2 asks for changes only")
check(max_out == 6144, f"smaller output budget: {max_out}")
check("equal opportunity" not in user_prompt and "401(k)" not in user_prompt and "Kafka" in user_prompt, "JD boilerplate trimmed, requirements kept")
new_docx = base64.b64decode(res["docx_base64"])
_, new_paras = docx_editor.load_paragraphs(new_docx)
texts = [p["text"] for p in new_paras]
check(any("Spark, Kafka, Python" in t for t in texts) and "Big Data: Spark, Kafka, Hadoop, Hive, Airflow" in texts, "changed paragraphs applied")
ins = texts.index("Streamed real-time transaction events with Kafka into Spark jobs for fraud reporting.")
check(texts[ins - 1] == paras[last_bullet]["text"], "[+N] bullet inserted right after paragraph N")
unchanged_old = [p["text"] for p in paras if p["index"] not in (summary_idx, skills_idx)]
check(all(t in texts for t in unchanged_old) and len(texts) == len(paras) + 1, "every other paragraph identical, nothing lost")
t = res.get("timings", {})
check(t.get("rewrite_mode") == "changes" and t.get("analysis_cached") is False and all(k in t for k in ("read_resume_s", "analysis_s", "rewrite_s", "apply_s", "total_s")),
      f"timings reported: {t}")

# ---- how much shorter the step-2 answer is (what the AI has to write)
full_answer = docx_editor.numbered_listing(paras)
saving = 100 * (1 - len(CHANGES) / len(full_answer))
print(f"step-2 answer on this {len(paras)}-paragraph resume: full rewrite {len(full_answer)} chars -> changes only {len(CHANGES)} chars ({saving:.0f}% shorter)")
check(saving > 70, "changes-only answer is much shorter")

# ---- Re-optimize: same resume + JD -> cached analysis, one AI call
calls.clear()
res2 = rb.optimize_resume_for_jd("", JD, docx_bytes=DOCX)
check(calls == ["rewrite"] and res2["timings"]["analysis_cached"] is True and res2.get("optimized"), f"re-optimize skips the analysis call: {calls}")
calls.clear()
rb.optimize_resume_for_jd("", JD.replace("- AWS experience required", "- AWS experience required\n- Snowflake experience required"), docx_bytes=DOCX)
check(calls == ["analysis", "rewrite"], "a different JD runs a fresh analysis")

# ---- plain-text path, changes-only
text_resume = docx_editor.extract_text(DOCX)
tp = rb._text_paragraphs(text_resume)
t_summary = next(p["index"] for p in tp if p["text"].startswith("Data Engineer with 8"))
CHANGES_TEXT = f"[{t_summary}] Data Engineer with 8 years of experience building batch and streaming pipelines on AWS using Spark, Kafka, Python and SQL.\n"
rb._generate = lambda c, ty, s, p, want_json, max_tokens: (json.dumps(ANALYSIS), False, "stub") if want_json else (CHANGES_TEXT, False, "stub")
rb._ANALYSIS_CACHE.clear()
res3 = rb.optimize_resume_for_jd(text_resume, JD)
check(res3.get("optimized") and "Spark, Kafka, Python and SQL" in res3["updated_resume_text"]
      and res3["updated_resume_text"].count("\n") == text_resume.count("\n"), f"text path: {res3.get('not_optimized_reason')}")

# ---- RESUME_REWRITE_MODE=full still works (the old behaviour)
rb.REWRITE_MODE = "full"
full_ans = "\n".join(f"[{p['index']}] {p['text'].replace('Spark, Python and SQL', 'Spark, Kafka, Python and SQL')}" for p in paras)
rb._generate = lambda c, ty, s, p, want_json, max_tokens: (json.dumps(ANALYSIS), False, "stub") if want_json else (full_ans, False, "stub")
rb._ANALYSIS_CACHE.clear()
res4 = rb.optimize_resume_for_jd("", JD, docx_bytes=DOCX)
check(res4.get("optimized") and res4["timings"]["rewrite_mode"] == "full", f"full mode: {res4.get('not_optimized_reason')}")
rb.REWRITE_MODE = "changes"

# ---- an answer that changes nothing / adds an unapproved technology is still refused
rb._generate = lambda c, ty, s, p, want_json, max_tokens: (json.dumps(ANALYSIS), False, "stub") if want_json else (f"[{skills_idx}] Big Data: Spark, Kubernetes, Hadoop, Hive, Airflow", False, "stub")
rb._ANALYSIS_CACHE.clear()
res5 = rb.optimize_resume_for_jd("", JD, docx_bytes=DOCX)
check(not res5.get("optimized") and "Kubernetes" in json.dumps(res5.get("skipped_edits", [])) + (res5.get("not_optimized_reason") or ""),
      f"unapproved technology refused: {res5.get('not_optimized_reason')}")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: optimizer speed - changes-only rewrite applied in place, [+N] inserts, analysis cache, JD trimmed, text path, full mode, safety checks, timings.")
