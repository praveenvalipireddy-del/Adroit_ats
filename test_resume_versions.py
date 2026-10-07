"""Saved JD-tailored resume versions (resume_versions.py, table optimized_resumes) and drafts using them.

Naming <Name>_<PrimarySkill>.docx with special characters -> "_", never overwritten (_v2, _v3);
primary skill = the matched skill the JD names, else the consultant's first skill. Saving: an
unedited version keeps the optimizer's in-place edited Word file; line edits are written into that
file (formatting kept); added lines rebuild a clean file and say so. Only the consultant's own
recruiter (or an admin) can save / list / download. Paste & Draft attaches the chosen version, a
1-Click Draft for a job attaches the latest version saved for that job, otherwise the original.
Gmail is never contacted (imaplib stubbed). Throwaway SQLite DB. Labelled test fixtures.
Run: python test_resume_versions.py
"""
import base64
import email
import email.policy
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_versions.db")
os.environ["DATABASE_URL"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import docx  # noqa: E402
import docx_editor  # noqa: E402
import gmail_multi_manager as gm  # noqa: E402
import models  # noqa: E402
import resume_versions as rv  # noqa: E402

failures, captured = [], []


def check(cond, msg):
    if not cond:
        failures.append(msg)


class FakeIMAP:
    def __init__(self, host, port=None):
        pass

    def login(self, user, pw):
        pass

    def append(self, folder, flags, when, raw):
        captured.append(email.message_from_bytes(raw, policy=email.policy.default))

    def logout(self):
        pass


gm.imaplib.IMAP4_SSL = FakeIMAP

# ---- naming
check(rv.make_filename("Koushik Versiontest", "Java Full Stack", []) == "Koushik_Versiontest_Java_Full_Stack.docx", "basic name")
check(rv.make_filename("Ravi O'Neil-K.", "C#/.NET Core", []) == "Ravi_O_Neil_K_C_NET_Core.docx", "special characters -> _")
taken = ["Koushik_Versiontest_Java_Full_Stack.docx", "koushik_versiontest_java_full_stack_v2.docx"]
check(rv.make_filename("Koushik Versiontest", "Java Full Stack", taken) == "Koushik_Versiontest_Java_Full_Stack_v3.docx", "never overwrite: _v3")
check(rv.primary_skill({"mandatory_matched_skills": ["Spring Boot", "Java Full Stack"]}, {}, "Need Java Full Stack with React") == "Java Full Stack",
      "the matched skill the JD names")
check(rv.primary_skill({"mandatory_matched_skills": []}, {"primary_skills": "Python, AWS"}, "x") == "Python", "fallback: profile skill")

# ---- fixtures
d = docx.Document()
for line in ["KOUSHIK VERSIONTEST", "PROFESSIONAL SUMMARY", "Java developer with 6 years of experience.", "TECHNICAL SKILLS", "Java, Spring Boot"]:
    d.add_paragraph(line)
d.paragraphs[0].runs[0].bold = True
buf = io.BytesIO()
d.save(buf)
ORIGINAL = buf.getvalue()
ai_doc = docx.Document(io.BytesIO(ORIGINAL))
ai_doc.paragraphs[4].runs[0].text = "Java, Spring Boot, React"
buf = io.BytesIO()
ai_doc.save(buf)
AI_DOCX = buf.getvalue()
AI_TEXT = docx_editor.extract_text(AI_DOCX)

uid = lambda u: u["id"] if isinstance(u, dict) else u  # noqa: E731
rec = {"id": uid(models.create_user("Version Rec", "ver-rec@example.invalid", "Ver-1", role="Recruiter")), "name": "Version Rec", "role": "Recruiter"}
other = {"id": uid(models.create_user("Version Other", "ver-other@example.invalid", "Ver-2", role="Recruiter")), "name": "Version Other", "role": "Recruiter"}
cid = models.create_candidate("Koushik Versiontest", "koushik-ver@example.invalid", title="Java Developer", primary_skills="Java, Spring Boot",
                              gmail_account="koushik-ver@example.invalid", assigned_user_id=rec["id"])
models.update_candidate(cid, gmail_app_password="test-app-password-not-real", country="United States")
models.save_resume_file(cid, "Koushik_original.docx", ORIGINAL)
client = app_module.app.test_client()


def login(u):
    with client.session_transaction() as s:
        s["user"] = u


def save(**kw):
    body = {"candidate_id": cid, "jd_text": "Java Full Stack developer: Java, Spring Boot, React required.", "ai_text": AI_TEXT,
            "edited_text": AI_TEXT, "docx_base64": base64.b64encode(AI_DOCX).decode(), "primary_skill": "Java Full Stack", **kw}
    return client.post("/api/optimized-resumes", json=body)


check(client.post("/api/optimized-resumes", json={}).status_code == 401, "login required")
login(other)
check(save().status_code == 404, "another recruiter cannot save for this consultant")
login(rec)

# unedited -> the optimizer's in-place edited file, byte for byte
r = save().get_json()
check(r.get("filename") == "Koushik_Versiontest_Java_Full_Stack.docx" and r["format_preserved"] and r["note"] == "", f"first save: {r}")
v1 = rv.get(r["id"])
check(v1["data"] == AI_DOCX, "unedited: the AI-edited Word file is stored as is")

# line edit -> written into the same Word file (bold name kept)
edited = AI_TEXT.replace("Java developer with 6 years of experience.", "Java full stack developer with 6 years of experience.")
r2 = save(edited_text=edited).get_json()
check(r2["filename"] == "Koushik_Versiontest_Java_Full_Stack_v2.docx" and r2["format_preserved"], f"second save -> _v2, formatting kept: {r2}")
d2 = docx.Document(io.BytesIO(rv.get(r2["id"])["data"]))
check(d2.paragraphs[2].text == "Java full stack developer with 6 years of experience." and d2.paragraphs[0].runs[0].bold, "edit applied, bold name kept")

# added line -> clean rebuild, with a note
r3 = save(edited_text=AI_TEXT + "\nCERTIFICATIONS\nOracle Certified Java Programmer").get_json()
check(r3["filename"].endswith("_v3.docx") and not r3["format_preserved"] and "rebuilt" in r3["note"], f"added lines -> rebuilt: {r3}")
check("Oracle Certified Java Programmer" in docx_editor.extract_text(rv.get(r3["id"])["data"]), "the added line is in the file")
check(save(edited_text="   ").status_code == 400, "empty text refused")
check(save(docx_base64="bm90IGEgZG9jeA==").status_code == 400, "a non-docx payload is refused")

# list + download + access
lst = client.get(f"/api/optimized-resumes?candidate_id={cid}").get_json()["versions"]
check([x["filename"] for x in lst][:3] == [r3["filename"], r2["filename"], r["filename"]] and "data" not in lst[0], "listed newest first, no file bytes")
dl = client.get(f"/api/optimized-resumes/{r2['id']}/download")
check(dl.status_code == 200 and "Koushik_Versiontest_Java_Full_Stack_v2.docx" in dl.headers.get("Content-Disposition", ""), "download with its name")
login(other)
check(client.get(f"/api/optimized-resumes/{r2['id']}/download").status_code == 404, "other recruiter can't download")
check(client.get(f"/api/optimized-resumes?candidate_id={cid}").status_code == 404, "other recruiter can't list")
login(rec)
check(models.get_resume_file(cid)["data"] == ORIGINAL, "the original resume is never touched")

# ---- Paste & Draft attaches the chosen version
r = client.post("/api/outreach/paste-and-draft", json={"candidate_id": cid, "raw_jd_text": "Java Full Stack developer needed",
                                                         "recruiter_email": "rec@example.invalid", "job_title": "Java Full Stack Developer",
                                                         "optimized_resume_id": r2["id"]}).get_json()
att = [p.get_filename() for p in captured[-1].iter_attachments()]
check(r.get("success") and att == ["Koushik_Versiontest_Java_Full_Stack_v2.docx"] and r.get("resume_filename") == att[0], f"paste draft attaches v2: {att}")
r = client.post("/api/outreach/paste-and-draft", json={"candidate_id": cid, "raw_jd_text": "x", "recruiter_email": "rec@example.invalid",
                                                         "job_title": "Java Developer"}).get_json()
check([p.get_filename() for p in captured[-1].iter_attachments()] == ["Koushik_original.docx"], "no version chosen -> the original")
cid2 = models.create_candidate("Other Versiontest", "other-ver@example.invalid", assigned_user_id=rec["id"])
check(client.post("/api/outreach/paste-and-draft", json={"candidate_id": cid2, "raw_jd_text": "x", "recruiter_email": "rec@example.invalid",
                                                           "job_title": "Java Developer", "optimized_resume_id": r2["id"]}).status_code == 400,
      "a version can't be attached to another consultant's draft")

# ---- 1-Click Draft for a job attaches the latest version saved for THAT job
jid = models.save_or_update_scraped_job({"title": "Java Full Stack Developer", "company": "Version Co", "source": "Dice",
                                         "url": "https://www.dice.com/job-detail/vertest", "recruiter_email": "rec2@example.invalid", "description": "Java"})
jid = jid["id"] if isinstance(jid, dict) else jid
client.post("/api/outreach/create-draft", json={"candidate_id": cid, "job_id": jid})
check([p.get_filename() for p in captured[-1].iter_attachments()] == ["Koushik_original.docx"], "no version for the job -> original")
rj = save(job_id=jid).get_json()
client.post("/api/outreach/create-draft", json={"candidate_id": cid, "job_id": jid})
check([p.get_filename() for p in captured[-1].iter_attachments()] == [rj["filename"]], f"job draft attaches the version saved for it: {rj['filename']}")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: resume versions - Name_Skill.docx naming + _vN, primary skill, save keeps formatting / applies edits / rebuilds, access, drafts attach the saved version.")
