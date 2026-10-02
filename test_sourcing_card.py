"""Sourcing tracker card (backend): contact + details fields, validation, owner rule, change history,
duplicate warnings, Add to Bench, admin reassign, quick filters, Excel contact columns admin-only.
Standalone script on a throwaway SQLite DB; no network. People are labelled test fixtures.
Run: python test_sourcing_card.py
"""
import io
import os
import sys
import tempfile
from datetime import date, timedelta

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_card.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def uid(u):
    return u["id"] if isinstance(u, dict) else u


ADMIN = {"id": uid(models.create_user("Card Admin", "card-admin@example.invalid", "Card-1", role="Admin")), "name": "Card Admin", "role": "Admin"}
PRAMOD = {"id": uid(models.create_user("Pramod Card", "card-pramod@example.invalid", "Card-2", role="Recruiter")), "name": "Pramod Card", "role": "Recruiter"}
RANI = {"id": uid(models.create_user("Rani Card", "card-rani@example.invalid", "Card-3", role="Recruiter")), "name": "Rani Card", "role": "Recruiter"}


def item(slug, first):
    return {"linkedinUrl": f"https://www.linkedin.com/in/{slug}.example.invalid", "firstName": first, "lastName": "Cardtest",
            "headline": "Data Engineer", "location": {"linkedinText": "Dallas, Texas"},
            "currentPosition": [{"position": "Data Engineer", "companyName": "Test Co"}],
            "education": [{"schoolName": "CBIT", "degree": "B.Tech", "endDate": {"year": 2018}},
                          {"schoolName": "UNT", "degree": "MS", "endDate": {"year": 2021}}]}


li.ingest_profiles(li.ApifyProfileProvider().to_profiles([item("card-a", "Ravi"), item("card-b", "Sita"), item("card-c", "Tara")]), "apify", None)
models.create_candidate("Bench Existing", "bench.exists@example.invalid", phone="+1 214 555 0111", assigned_user_id=RANI["id"])
client = app_module.app.test_client()


def login(u):
    with client.session_transaction() as s:
        s["user"] = {**u, "email": "x@example.invalid"}


def rows(**q):
    r = client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-cbit", **q})
    return {x["name"]: x for x in (r.get_json() or {}).get("results", [])}


login(PRAMOD)
ids = {n: r["id"] for n, r in rows().items()}
ravi, sita, tara = ids["Ravi Cardtest"], ids["Sita Cardtest"], ids["Tara Cardtest"]

# ---- card defaults: location from LinkedIn, everything else empty, no owner
c = client.get(f"/api/sourcing/profiles/{ravi}/card").get_json()
check(c["fields"]["current_location"] == "Dallas, Texas" and c["fields"]["contact_email"] == "" and c["fields"]["visa_status"] == ""
      and c["owner_user_id"] is None and c["status"] == "New", f"card defaults: {c['fields']} owner={c['owner_user_id']}")
check("CBIT" in "".join(e["institution_name"] for e in c["education"]), "card shows education")
check(c["options"]["statuses"][2] == "No response" and "H1B" in c["options"]["visa"], "options")

# ---- validation
bad = [({"contact_email": "not-an-email"}, "email"), ({"contact_phone": "4695550100"}, "country code"),
       ({"contact_phone": "+1 23"}, "complete"), ({"visa_status": "Martian"}, "visa"), ({"open_to_relocate": "Maybe"}, "relocate"),
       ({"availability": "someday"}, "availability"), ({"follow_up_date": "2026-13-40"}, "date")]
for body, word in bad:
    r = client.post(f"/api/sourcing/profiles/{ravi}/card", json=body)
    check(r.status_code == 400 and word in r.get_json().get("error", "").lower(), f"{body} -> {r.status_code} {r.get_json()}")

# ---- save fields: normalized, history records WHICH fields, no owner just for saving details
follow = (date.today() + timedelta(days=3)).isoformat()
r = client.post(f"/api/sourcing/profiles/{ravi}/card", json={
    "contact_email": "Ravi.K@Example.invalid", "contact_phone": "+1 (469) 555-0100", "visa_status": "H1B",
    "current_location": "Irving, TX", "open_to_relocate": "Yes", "expected_rate": "$65/hr C2C", "availability": "2 weeks",
    "follow_up_date": follow})
d = r.get_json()
check(r.status_code == 200 and d["fields"]["contact_email"] == "ravi.k@example.invalid" and d["fields"]["contact_phone"] == "+14695550100"
      and d["warnings"] == [], f"save: {r.status_code} {d.get('fields')} {d.get('warnings')}")
hist = [t for t in d["thread"] if t["kind"] == "field"]
check(hist and hist[0]["author"] == "Pramod Card" and "email" in hist[0]["text"] and "phone" in hist[0]["text"]
      and "follow-up date" in hist[0]["text"], f"field history: {hist[:1]}")
n = len(d["thread"])
d2 = client.post(f"/api/sourcing/profiles/{ravi}/card", json={"visa_status": "H1B"}).get_json()
check(len(d2["thread"]) == n, "saving unchanged values adds no history")
check(d["owner_user_id"] is None, "saving details alone doesn't claim ownership")

# ---- owner rule: first recruiter to move past New owns the candidate; later ones don't take over
client.post(f"/api/sourcing/profiles/{ravi}/status", json={"status": "Contacted"})
login(RANI)
d = client.post(f"/api/sourcing/profiles/{ravi}/status", json={"status": "No response"}).get_json()
check(d["owner_name"] == "Pramod Card" and d["status"] == "No response", f"owner stays Pramod: {d['owner_name']}")
check(client.post(f"/api/sourcing/profiles/{ravi}/owner", json={"user_id": RANI["id"]}).status_code == 403, "recruiter cannot reassign")

# ---- duplicates: same email on another sourced candidate / phone of a bench consultant
d = client.post(f"/api/sourcing/profiles/{sita}/card", json={"contact_email": "ravi.k@example.invalid", "contact_phone": "+1 214-555-0111"}).get_json()
w = " | ".join(d.get("warnings", []))
check("Email also saved on sourced candidate Ravi Cardtest (owner Pramod Card)" in w and "Phone matches bench consultant Bench Existing" in w,
      f"duplicate warnings: {w}")
d = client.post(f"/api/sourcing/profiles/{sita}/card", json={"contact_email": "bench.exists@example.invalid", "contact_phone": ""}).get_json()
check(any("Already on the bench as Bench Existing (Rani Card)" in x for x in d.get("warnings", [])), f"bench email warning: {d.get('warnings')}")

# ---- admin reassign
login(ADMIN)
d = client.post(f"/api/sourcing/profiles/{ravi}/owner", json={"user_id": RANI["id"]}).get_json()
check(d["owner_name"] == "Rani Card" and any(t["kind"] == "owner" and "Rani Card" in t["text"] and "Card Admin" in t["text"] for t in d["thread"]),
      "admin reassign recorded")

# ---- Add to Bench
login(PRAMOD)
r = client.post(f"/api/sourcing/profiles/{tara}/add-to-bench", json={})
check(r.status_code == 400 and "email" in r.get_json()["error"], "Add to Bench needs an email")
client.post(f"/api/sourcing/profiles/{tara}/card", json={"contact_email": "tara@example.invalid", "contact_phone": "+1 972 555 0199",
                                                        "visa_status": "OPT", "expected_rate": "$55/hr"})
r = client.post(f"/api/sourcing/profiles/{tara}/add-to-bench", json={})
d = r.get_json()
check(r.status_code == 200 and d["status"] == "Added to bench" and d["bench_candidate_id"] and d["owner_name"] == "Pramod Card",
      f"add to bench: {r.status_code} {d.get('status')} {d.get('owner_name')}")
bench = models.get_candidate_by_id(d["candidate_id"])
check(bench and bench["email"] == "tara@example.invalid" and bench["phone"] == "+19725550199" and bench["visa_status"] == "OPT"
      and bench["target_rate"] == "$55/hr" and bench["assigned_user_id"] == PRAMOD["id"] and "linkedin.com/in/card-c" in (bench["resume_summary"] or ""),
      f"bench consultant: {bench}")
check(client.post(f"/api/sourcing/profiles/{tara}/add-to-bench", json={}).status_code == 400, "no second bench copy")

# ---- quick filters
check(sorted(rows(mine="1")) == ["Tara Cardtest"], f"My candidates (Pramod): {sorted(rows(mine='1'))}")
client.post(f"/api/sourcing/profiles/{sita}/card", json={"follow_up_date": (date.today() - timedelta(days=1)).isoformat()})
check(sorted(rows(followup_due="1")) == ["Sita Cardtest"], f"follow-up due: {sorted(rows(followup_due='1'))}")
check(sorted(rows(has_contact="1")) == ["Ravi Cardtest", "Sita Cardtest", "Tara Cardtest"], "has contact")
row = rows()["Ravi Cardtest"]
check(row["has_email"] and row["has_phone"] and row["visa"] == "H1B" and row["owner_name"] == "Rani Card" and row["follow_up_date"] == follow
      and "contact_email" not in row and "contact_phone" not in row, f"table row: {row}")

# ---- Excel: email + phone only for admins
def excel_header(u):
    login(u)
    r = client.get("/api/education/export-xlsx", query_string={"filter": "A", "institution_id": "in-cbit"})
    ws = load_workbook(io.BytesIO(r.data))["Candidates"]
    header = [x.value for x in ws[1]]
    return header, ws


h, ws = excel_header(PRAMOD)
check("Email" not in h and "Phone" not in h and "Visa" in h and "Owner" in h and "Next Follow-up" in h, f"recruiter Excel: {h}")
h, ws = excel_header(ADMIN)
check("Email" in h and "Phone" in h, f"admin Excel: {h}")
vals = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=h.index("Email") + 1).value for i in range(2, ws.max_row + 1)}
check(vals.get("Ravi Cardtest") == "ravi.k@example.invalid", f"admin Excel email: {vals}")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: tracker card - fields/validation, history, owner rule, duplicates, admin reassign, Add to Bench, quick filters, Excel contacts admin-only.")
