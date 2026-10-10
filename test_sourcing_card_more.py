"""Sourcing card - more details: GitHub / portfolio links, experience, primary skills, work mode,
preferred locations, current rate, visa expiry, certifications, referred by, marital status, and an
attached resume (Word / PDF, 5 MB).

Values are validated (GitHub must be a github.com link, a link without https:// gets it, experience
is a number of years, dates are real dates, dropdowns only take their options); the activity history
names what changed; Add to Bench copies skills, experience, links and the resume to the consultant -
but never the marital status. Throwaway SQLite DB, no network. People are labelled test fixtures.
Run: python test_sourcing_card_more.py
"""
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_card_more.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import docx  # noqa: E402
import app as app_module  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import models  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


uid = lambda u: u["id"] if isinstance(u, dict) else u  # noqa: E731
REC = {"id": uid(models.create_user("Kiran Moretest", "more-rec@example.invalid", "More-1", role="Recruiter")), "name": "Kiran Moretest", "role": "Recruiter"}
li.ingest_profiles(li.ApifyProfileProvider().to_profiles([{
    "linkedinUrl": "https://www.linkedin.com/in/more-1.example.invalid", "firstName": "Anil", "lastName": "Moretest",
    "headline": "Java Developer", "location": {"linkedinText": "Dallas, Texas"}, "currentPosition": [],
    "education": [{"schoolName": "JNTUH", "degree": "B.Tech", "endDate": {"year": 2018}}]}]), "apify", None)
conn = models.get_db_connection()
cur = conn.cursor()
cur.execute("SELECT id FROM linkedin_profiles WHERE linkedin_url LIKE '%more-1%'")
pid = cur.fetchone()["id"]
conn.close()

client = app_module.app.test_client()
with client.session_transaction() as s:
    s["user"] = {**REC, "email": "more-rec@example.invalid"}
URL = f"/api/sourcing/profiles/{pid}/card"

c = client.get(URL).get_json()
for k in ("github_url", "portfolio_url", "experience_years", "primary_skills", "work_mode", "preferred_locations",
          "current_rate", "visa_expiry", "certifications", "referred_by", "marital_status"):
    check(k in c["fields"] and c["fields"][k] == "", f"new field starts empty: {k}")
check(c["options"]["work_mode"] == ["Remote", "Hybrid", "Onsite", "Any"] and "Married" in c["options"]["marital"], "dropdown options")
check(c["resume"] is None, "no resume yet")

# ---- validation
for bad, why in (({"github_url": "https://gitlab.com/anil"}, "GitHub"), ({"portfolio_url": "not a link"}, "Portfolio"),
                 ({"experience_years": "eight"}, "Experience"), ({"experience_years": "75"}, "between"),
                 ({"visa_expiry": "2027-02-30"}, "real date"), ({"work_mode": "Mars"}, "work mode"),
                 ({"marital_status": "Complicated"}, "marital"), ({"primary_skills": "x" * 501}, "too long")):
    r = client.post(URL, json=bad)
    check(r.status_code == 400 and why.lower() in r.get_json()["error"].lower(), f"rejects {bad}: {r.get_json()}")

# ---- save everything
r = client.post(URL, json={"contact_email": "anil.more@example.invalid", "github_url": "github.com/anil-more",
                           "portfolio_url": "anil-more.example.dev", "experience_years": "7.50", "primary_skills": "Java, Spring Boot, AWS",
                           "work_mode": "Hybrid", "preferred_locations": "Dallas, Austin", "current_rate": "$55/hr W2",
                           "visa_expiry": "2027-06-30", "certifications": "AWS Solutions Architect", "referred_by": "Ravi K",
                           "marital_status": "Married"})
f = r.get_json()["fields"]
check(r.status_code == 200, f"saved: {r.get_json()}")
check(f["github_url"] == "https://github.com/anil-more" and f["portfolio_url"] == "https://anil-more.example.dev", f"links get https://: {f['github_url']} {f['portfolio_url']}")
check(f["experience_years"] == "7.5" and f["work_mode"] == "Hybrid" and f["marital_status"] == "Married" and f["visa_expiry"] == "2027-06-30", "values stored")
hist = " ".join(t["text"] for t in r.get_json()["thread"])
check("GitHub" in hist and "primary skills" in hist and "visa expiry" in hist, f"history names the changes: {hist}")

# ---- resume
d = docx.Document()
d.add_paragraph("ANIL MORETEST")
d.add_paragraph("Java developer with Spring Boot and AWS.")
buf = io.BytesIO()
d.save(buf)
DOCX = buf.getvalue()
r = client.post(f"/api/sourcing/profiles/{pid}/resume", data={"file": (io.BytesIO(b"hello"), "notes.txt")}, content_type="multipart/form-data")
check(r.status_code == 400 and "Word" in r.get_json()["error"], "only .docx / .pdf")
r = client.post(f"/api/sourcing/profiles/{pid}/resume", data={"file": (io.BytesIO(b"x" * (5 * 1024 * 1024 + 10)), "big.pdf")}, content_type="multipart/form-data")
check(r.status_code == 400 and "5 MB" in r.get_json()["error"], "5 MB limit")
r = client.post(f"/api/sourcing/profiles/{pid}/resume", data={"file": (io.BytesIO(DOCX), "Anil Resume.docx")}, content_type="multipart/form-data")
check(r.status_code == 200 and r.get_json()["resume"]["filename"] == "Anil_Resume.docx" and r.get_json()["resume"]["uploaded_by"] == "Kiran Moretest",
      f"resume attached: {r.get_json().get('resume')}")
dl = client.get(f"/api/sourcing/profiles/{pid}/resume")
check(dl.status_code == 200 and dl.data == DOCX and "Anil_Resume.docx" in dl.headers.get("Content-Disposition", ""), "resume downloads unchanged")

# ---- Add to Bench: skills, experience, links, resume carried; marital status NOT
r = client.post(f"/api/sourcing/profiles/{pid}/add-to-bench")
check(r.status_code == 200, f"added to bench: {r.get_json()}")
cand = models.get_candidate_by_id(r.get_json()["candidate_id"], user_id=REC["id"], is_admin=True)
check(cand["primary_skills"] == "Java, Spring Boot, AWS" and cand["experience_years"] in (8, "8"), f"skills + experience: {cand['primary_skills']} {cand['experience_years']}")
check("https://github.com/anil-more" in cand["resume_summary"] and "anil-more.example.dev" in cand["resume_summary"], f"links on the consultant: {cand['resume_summary']}")
rf = models.get_resume_file(cand["id"])
check(rf and bytes(rf["data"]) == DOCX and cand["resume_filename"] == "Anil_Resume.docx" and "Spring Boot" in (cand.get("resume_text") or ""), "resume copied to the consultant")
check("married" not in str(cand).lower(), "marital status never copied to the consultant")

# ---- remove resume
r = client.delete(f"/api/sourcing/profiles/{pid}/resume")
check(r.status_code == 200 and r.get_json()["resume"] is None and client.get(f"/api/sourcing/profiles/{pid}/resume").status_code == 404, "resume removed")
with client.session_transaction() as s:
    s.clear()
check(client.get(f"/api/sourcing/profiles/{pid}/resume").status_code == 401, "login required")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: sourcing card more details - links, work details, extras, marital (card only), resume upload/download/remove, validation, Add to Bench carries skills/experience/links/resume.")
