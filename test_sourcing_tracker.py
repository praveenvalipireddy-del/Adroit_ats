"""Sourcing tracker: per-candidate status + comments (team-wide), status filter, Excel columns,
delete permissions. Standalone script on a throwaway SQLite DB; no network. People are labelled
test fixtures.
Run: python test_sourcing_tracker.py
"""
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_tracker.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def uid(u):
    return u["id"] if isinstance(u, dict) else u


ADMIN = {"id": uid(models.create_user("Trk Admin", "trk-admin@example.invalid", "Trk-1", role="Admin")), "name": "Trk Admin", "role": "Admin"}
PRAMOD = {"id": uid(models.create_user("Pramod Trk", "trk-pramod@example.invalid", "Trk-2", role="Recruiter")), "name": "Pramod Trk", "role": "Recruiter"}
OTHER = {"id": uid(models.create_user("Other Trk", "trk-other@example.invalid", "Trk-3", role="Recruiter")), "name": "Other Trk", "role": "Recruiter"}


def item(slug, first, college, year):
    return {"linkedinUrl": f"https://www.linkedin.com/in/{slug}.example.invalid", "firstName": first, "lastName": "Trktest",
            "headline": "Engineer", "location": {"linkedinText": "Dallas, Texas"}, "currentPosition": [],
            "education": [{"schoolName": college, "degree": "B.Tech", "endDate": {"year": year}},
                          {"schoolName": "UT Dallas", "degree": "MS", "endDate": {"year": year + 3}}]}


li.ingest_profiles(li.ApifyProfileProvider().to_profiles([item("trk-a", "Asha", "JNTUH", 2019), item("trk-b", "Bharat", "JNTUH", 2019)]),
                   "apify", ADMIN["id"])
client = app_module.app.test_client()


def login(u):
    with client.session_transaction() as s:
        s["user"] = {**u, "email": "x@example.invalid"}


def results(**q):
    r = client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-jntu-hyderabad", **q})
    return {x["name"]: x for x in (r.get_json() or {}).get("results", [])}


check(client.get("/api/sourcing/profiles/1/comments").status_code == 401, "tracker requires login")
login(PRAMOD)
rows = results()
asha, bharat = rows["Asha Trktest"], rows["Bharat Trktest"]
check(asha["status"] == "New" and asha["comment_count"] == 0 and asha["latest_comment"] is None, f"defaults: {asha}")

# ---- status
r = client.post(f"/api/sourcing/profiles/{asha['id']}/status", json={"status": "Contacted"})
check(r.status_code == 200 and r.get_json()["status"] == "Contacted", f"set status: {r.get_json()}")
thread = r.get_json()["thread"]
st = next((c for c in thread if c["kind"] == "status"), None)
check(st and st["text"] == "Status: New → Contacted" and st["author"] == "Pramod Trk", f"status change recorded in the thread: {thread[:2]}")
check(any(c["kind"] == "owner" and c["text"] == "Owner: Pramod Trk" for c in thread), "first recruiter to contact becomes owner")
check(client.post(f"/api/sourcing/profiles/{asha['id']}/status", json={"status": "Hired!!"}).status_code == 400, "unknown status rejected")
check(client.post("/api/sourcing/profiles/999999/status", json={"status": "Contacted"}).status_code == 404, "unknown candidate")
n_before = len(client.get(f"/api/sourcing/profiles/{asha['id']}/comments").get_json()["thread"])
client.post(f"/api/sourcing/profiles/{asha['id']}/status", json={"status": "Contacted"})
check(len(client.get(f"/api/sourcing/profiles/{asha['id']}/comments").get_json()["thread"]) == n_before, "same status again adds no history")

# ---- comments
check(client.post(f"/api/sourcing/profiles/{asha['id']}/comments", json={"comment": "   "}).status_code == 400, "empty comment rejected")
check(client.post(f"/api/sourcing/profiles/{asha['id']}/comments", json={"comment": "x" * 2001}).status_code == 400, "too long rejected")
r = client.post(f"/api/sourcing/profiles/{asha['id']}/comments", json={"comment": "Called - interested, expects $60/hr, H1B"})
check(r.status_code == 200, f"add comment: {r.status_code}")
login(OTHER)
client.post(f"/api/sourcing/profiles/{asha['id']}/comments", json={"comment": "Sent resume to client <script>x</script>"})
thread = client.get(f"/api/sourcing/profiles/{asha['id']}/comments").get_json()["thread"]
comments = [c for c in thread if c["kind"] == "comment"]
check([c["author"] for c in comments] == ["Other Trk", "Pramod Trk"], f"team-wide thread, newest first: {[c['author'] for c in comments]}")

rows = results()
check(rows["Asha Trktest"]["comment_count"] == 2 and rows["Asha Trktest"]["latest_comment"]["author"] == "Other Trk"
      and rows["Asha Trktest"]["status"] == "Contacted", f"row shows count/latest/status: {rows['Asha Trktest']}")
check(rows["Bharat Trktest"]["comment_count"] == 0, "other candidates unaffected")

# ---- delete: only the author or an admin; status records can't be deleted
pramod_comment = next(c for c in comments if c["author"] == "Pramod Trk")
check(client.delete(f"/api/sourcing/comments/{pramod_comment['id']}").status_code == 403, "another recruiter cannot delete")
status_rec = next(c for c in thread if c["kind"] == "status")
login(ADMIN)
check(client.delete(f"/api/sourcing/comments/{status_rec['id']}").status_code == 404, "status history cannot be deleted")
r = client.delete(f"/api/sourcing/comments/{pramod_comment['id']}")
check(r.status_code == 200 and all(c["id"] != pramod_comment["id"] for c in r.get_json()["thread"]), "admin can delete")
login(OTHER)
own = next(c for c in comments if c["author"] == "Other Trk")
check(client.delete(f"/api/sourcing/comments/{own['id']}").status_code == 200, "author can delete own comment")
client.post(f"/api/sourcing/profiles/{asha['id']}/comments", json={"comment": "Interested - follow up Monday"})

# ---- status filter
client.post(f"/api/sourcing/profiles/{bharat['id']}/status", json={"status": "Interested"})
check(list(results(status="Contacted")) == ["Asha Trktest"], f"filter Contacted: {list(results(status='Contacted'))}")
check(list(results(status="Interested")) == ["Bharat Trktest"], "filter Interested")
check(list(results(status="New")) == [], "nobody New any more")
check(client.get("/api/education/search", query_string={"filter": "A", "institution_id": "in-jntu-hyderabad", "status": "Bogus"}).status_code == 400,
      "unknown status filter rejected")
check(sorted(results()) == ["Asha Trktest", "Bharat Trktest"], "no status filter = everyone")

# ---- Excel
r = client.get("/api/education/export-xlsx", query_string={"filter": "A", "institution_id": "in-jntu-hyderabad"})
ws = load_workbook(io.BytesIO(r.data))["Candidates"]
header = [c.value for c in ws[1]]
for col in ("Status", "Owner", "Latest Comment", "Comments"):
    check(col in header, f"Excel tracker column {col!r} missing: {header}")
rowmap = {ws.cell(row=i, column=1).value: {header[j - 1]: ws.cell(row=i, column=j).value for j in range(1, len(header) + 1)}
          for i in range(2, ws.max_row + 1)}
asha_x = rowmap.get("Asha Trktest") or {}
check(asha_x.get("Status") == "Contacted" and str(asha_x.get("Latest Comment")).startswith("Interested - follow up Monday (Other Trk,")
      and asha_x.get("Comments") == 1, f"Excel row: {asha_x}")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: sourcing tracker - status (with history), team comments, delete rules, status filter, Excel columns.")
