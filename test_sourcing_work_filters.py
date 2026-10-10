"""Sourcing: Open to work / Contract / Full-time / Remote / Hybrid / On-site filters + scraped counts.

The scraper keeps LinkedIn's openToWork flag and the current job's employmentType / workplaceType
(Full mode already returns them); a past job doesn't count; a blank value stays blank (never
guessed); a re-scrape replaces old values; older profiles without the data never match a work
checkbox and are counted as "no work data". Ticked boxes in one group are OR, groups AND. The
search returns total-scraped / this-week / per-checkbox counts; the Excel has the new columns.
Throwaway SQLite DB, no network. People are labelled test fixtures.
Run: python test_sourcing_work_filters.py
"""
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_work_filters.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import linkedin_sourcing  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


check("openToWork" in linkedin_sourcing._PROFILE_FIELDS and "experience" in linkedin_sourcing._PROFILE_FIELDS, "scraper keeps openToWork + experience")

# ---- parsing
cur_job = {"companyName": "Acme", "position": "Engineer", "employmentType": "Contract", "workplaceType": "Remote", "endDate": {"text": "Present"}}
old_job = {"companyName": "OldCo", "employmentType": "Full-time", "workplaceType": "On-site", "endDate": {"month": 3, "year": 2022, "text": "Mar 2022"}}
w = li.work_info({"openToWork": True, "currentPosition": [{"companyName": "Acme"}], "experience": [old_job, cur_job]})
check(w == {"open_to_work": 1, "current_employment_type": "Contract", "current_workplace_type": "Remote"}, f"current job, not the past one: {w}")
check(li.work_info({"linkedinUrl": "x"}) is None, "no fields -> no data (old scrape)")
w = li.work_info({"openToWork": False, "experience": [{"companyName": "Z", "endDate": None}]})
check(w == {"open_to_work": 0, "current_employment_type": "", "current_workplace_type": ""}, f"blank stays blank: {w}")

# ---- fixtures
ADMIN = {"id": models.create_user("Work Admin", "work-admin@example.invalid", "Work-1", role="Admin"), "name": "Work Admin", "role": "Admin"}
ADMIN["id"] = ADMIN["id"]["id"] if isinstance(ADMIN["id"], dict) else ADMIN["id"]


def item(slug, first, otw=None, emp=None, wp=None, with_work=True):
    it = {"linkedinUrl": f"https://www.linkedin.com/in/{slug}.example.invalid", "firstName": first, "lastName": "Worktest",
          "headline": "Engineer", "location": {"linkedinText": "Dallas, Texas"}, "currentPosition": [{"companyName": "Co"}],
          "education": [{"schoolName": "JNTUH", "degree": "B.Tech", "endDate": {"year": 2019}},
                        {"schoolName": "University of Texas at Dallas", "degree": "MS", "endDate": {"year": 2022}}]}
    if with_work:
        it["openToWork"] = otw
        it["experience"] = [{"companyName": "Co", "employmentType": emp, "workplaceType": wp, "endDate": {"text": "Present"}}]
    return it


li.ingest_profiles(li.ApifyProfileProvider().to_profiles([
    item("w-1", "Asha", True, "Contract", "Remote"), item("w-2", "Bala", False, "Full-time", "Hybrid"),
    item("w-3", "Chitra", True, "Full-time", "On-site"), item("w-4", "Deepak", False, "Freelance", None),
    item("w-5", "Esha", with_work=False)]), "apify", ADMIN["id"])

client = app_module.app.test_client()
with client.session_transaction() as s:
    s["user"] = {**ADMIN, "email": "work-admin@example.invalid"}


def names(**flags):
    r = client.get("/api/education/search", query_string={"filter": "P", "passout_year": "All", **flags})
    d = r.get_json()
    return sorted(x["name"].split()[0] for x in d.get("results", [])), d


all_names, d = names()
check(len(all_names) == 5, f"all five match the passout search: {all_names} {d.get('error')}")
c = d.get("counts", {})
check(c.get("total_scraped") == 5 and c.get("added_this_week") == 5 and c.get("matching") == 5, f"totals: {c}")
check((c.get("open_to_work"), c.get("emp_contract"), c.get("emp_fulltime"), c.get("wp_remote"), c.get("wp_hybrid"), c.get("wp_onsite"), c.get("no_work_data"))
      == (2, 2, 2, 1, 1, 1, 1), f"checkbox counts: {c}")
check(names(open_to_work=1)[0] == ["Asha", "Chitra"], "Open to work")
check(names(emp_contract=1)[0] == ["Asha", "Deepak"], "Contract includes Freelance")
check(names(emp_contract=1, emp_fulltime=1)[0] == ["Asha", "Bala", "Chitra", "Deepak"], "same group = OR, no-data profile excluded")
check(names(open_to_work=1, emp_fulltime=1)[0] == ["Chitra"], "groups AND")
check(names(wp_remote=1, wp_hybrid=1)[0] == ["Asha", "Bala"], "Remote + Hybrid")
n, d = names(open_to_work=1)
row = next(x for x in d["results"] if x["name"].startswith("Asha"))
check(row["open_to_work"] == 1 and row["employment_type"] == "Contract" and row["workplace_type"] == "Remote", f"row badges data: {row}")
check(d["counts"]["matching"] == 5 and d["total"] == 2, "counts stay for the whole search; total is after the checkboxes")

# ---- re-scrape replaces stale values
li.ingest_profiles(li.ApifyProfileProvider().to_profiles([item("w-1", "Asha", False, "Full-time", None)]), "apify", ADMIN["id"])
check(names(open_to_work=1)[0] == ["Chitra"] and "Asha" in names(emp_fulltime=1)[0] and "Asha" not in names(wp_remote=1)[0], "re-scrape updates the values")

# ---- Excel
x = client.get("/api/education/export-xlsx", query_string={"filter": "P", "passout_year": "All", "emp_fulltime": "1"})
ws = load_workbook(io.BytesIO(x.data)).active
hdr = [c.value for c in ws[1]]
check(x.status_code == 200 and {"Open to Work", "Employment Type", "Workplace"} <= set(hdr), f"Excel columns: {hdr}")
check(ws.max_row - 1 == 3, f"Excel follows the checkboxes: {ws.max_row - 1} rows")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: sourcing LinkedIn work filters - open to work / contract / full-time / remote / hybrid / on-site, OR within group, AND across, counts + totals, no-data excluded, re-scrape refresh, Excel columns.")
