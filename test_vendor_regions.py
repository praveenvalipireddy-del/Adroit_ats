"""Vendors: separate USA and India lists.

1. guess_region(): +91 / 91XXXXXXXXXX / '98765 43210' / .in email -> India; +1 / '(949) 688-0755' -> USA;
   a bare 10-digit number says nothing.
2. backfill_regions(): contacts without a region are sorted once (phone / email, else the recruiter's team).
3. Lists are separate: the same email may be in both; a duplicate in the same list is refused; moving a
   contact into a list that already has that email is refused; "looks like an India contact" hint.
4. Drafts only use the consultant's list: Paste-window match, Paste & Draft BCC, the Jobs-feed vendor tag.
5. GET /api/vendors?region=..., counts per list, Excel download of one list.
6. Edge: USA / India tabs with counts, switching lists, "Move to India", Add contact into the open list.
Gmail is never contacted (imaplib stub). Standalone; throwaway SQLite DB; labelled test fixtures.
Run: python test_vendor_regions.py
"""
import email
import io
import os
import socket
import sys
import tempfile
import threading

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_vendor_regions.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import gmail_multi_manager as gm  # noqa: E402
import india_job_scrapers  # noqa: E402
import models  # noqa: E402
import us_job_scrapers  # noqa: E402
import vendors  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402
from werkzeug.serving import make_server  # noqa: E402

us_job_scrapers.run_multi_source_us_scrape = lambda *a, **k: {"jobs": [], "count": 0}
india_job_scrapers.run_multi_source_india_scrape = lambda *a, **k: {"jobs": [], "count": 0}
failures, captured = [], []


def check(cond, msg):
    if not cond:
        failures.append(msg)


class FakeIMAP:
    def __init__(self, host, port=None):
        pass

    def login(self, user, pw):
        pass

    def append(self, folder, flags, when, raw):
        captured.append(raw)

    def logout(self):
        pass


gm.imaplib.IMAP4_SSL = FakeIMAP


def uid(u):
    return u["id"] if isinstance(u, dict) else u


# ---- 1. guessing a contact's region
for phone, mail, want in [("+91 98765 43210", "", "India"), ("919876543210", "", "India"), ("98765 43210", "", "India"),
                          ("", "hr@tcs.co.in", "India"), ("(949) 688-0755", "", "USA"), ("(949)688-0755", "", "USA"),
                          ("949-688-0755", "", "USA"), ("+1 949 688 0755", "", "USA"), ("19496880755", "", "USA"),
                          ("9496880755", "a@x.com", ""), ("", "a@x.com", "")]:
    check(vendors.guess_region(phone, mail) == want, f"guess_region({phone!r}, {mail!r}) = {vendors.guess_region(phone, mail)!r}, want {want!r}")

EMAIL, PASSWORD = "vreg-admin@example.invalid", "VReg-Admin-1"
ADMIN = {"id": uid(models.create_user("Vreg Admin", EMAIL, PASSWORD, role="Admin")), "name": "Vreg Admin", "role": "Admin"}
IN_REC = uid(models.create_user("Vreg Indiarec", "vreg-in@example.invalid", "VReg-2", role="Recruiter", region="India"))
US_REC = {"id": uid(models.create_user("Vreg Usrec", "vreg-us@example.invalid", "VReg-3", role="Recruiter")), "name": "Vreg Usrec", "role": "Recruiter"}

# ---- 2. one-time sorting of contacts that have no region yet (an existing database)
conn = models.get_db_connection()
cur = conn.cursor()
for owner, mail, phone in [(ADMIN["id"], "old-in-phone@oldco.example", "+91 90000 00001"), (ADMIN["id"], "old-us-phone@oldco.example", "(212) 555-0101"),
                           (ADMIN["id"], "old@oldco.co.in", ""), (IN_REC, "plain-in@oldco.example", ""), (ADMIN["id"], "plain-us@oldco.example", "")]:
    cur.execute("""INSERT INTO vendor_contacts (owner_user_id, company_name, company_norm, email, email_domain, phone, status, region)
                   VALUES (?, 'Oldco Test', 'oldco test', ?, ?, ?, 'active', NULL)""", (owner, mail, vendors.company_domain(mail), phone))
conn.commit()
n = vendors.backfill_regions(conn)
cur.execute("SELECT email, region FROM vendor_contacts")
got = {r[0]: r[1] for r in cur.fetchall()}
check(n == 5, f"5 contacts sorted: {n}")
check(got == {"old-in-phone@oldco.example": "India", "old-us-phone@oldco.example": "USA", "old@oldco.co.in": "India",
              "plain-in@oldco.example": "India", "plain-us@oldco.example": "USA"}, f"backfill: {got}")
check(vendors.backfill_regions(conn) == 0, "backfill runs once")
cur.execute("DELETE FROM vendor_contacts")
conn.commit()

# ---- 3. two lists for one recruiter
us_cid = models.create_candidate("Usha Vregtest", "usha-vreg@example.invalid", title="Java Developer", gmail_account="usha-vreg@example.invalid",
                                 assigned_user_id=ADMIN["id"])
in_cid = models.create_candidate("Indu Vregtest", "indu-vreg@example.invalid", title="Java Developer", gmail_account="indu-vreg@example.invalid",
                                 assigned_user_id=ADMIN["id"], country="India")
for c in (us_cid, in_cid):
    models.update_candidate(c, gmail_app_password="test-app-password-not-real")
vendors.import_rows(conn, ADMIN["id"], [{"company": "Globex Test", "name": "Us One", "email": "us1@globex.example"},
                                        {"company": "Globex Test", "name": "Both Lists", "email": "both@globex.example"},
                                        {"company": "Globex Test", "name": "Indian Phone", "email": "inph@globex.example", "phone": "+91 98765 43210"}],
                    region="USA")
r_in = vendors.import_rows(conn, ADMIN["id"], [{"company": "Globex Test", "name": "In One", "email": "in1@globex.example"},
                                               {"company": "Globex Test", "name": "Both Lists", "email": "both@globex.example"}], region="India")
check(r_in["added"] == 2, f"the same email may be in both lists: {r_in}")
check(vendors.import_rows(conn, ADMIN["id"], [{"company": "Globex Test", "email": "in1@globex.example"}], region="India")["duplicates"] == 1,
      "a duplicate in the same list is refused")
cur.execute("SELECT id, email, region FROM vendor_contacts")
ID = {(r[1], r[2]): r[0] for r in cur.fetchall()}
conn.close()

client = app_module.app.test_client()
with client.session_transaction() as s:
    s["user"] = ADMIN

lst = client.get("/api/vendors?region=India").get_json()
check(sorted(c["email"] for c in lst["contacts"]) == ["both@globex.example", "in1@globex.example"] and lst["region"] == "India",
      f"India list: {lst['contacts']}")
check(lst["region_counts"] == {"USA": 3, "India": 2}, f"counts: {lst['region_counts']}")
us_list = client.get("/api/vendors?region=USA").get_json()["contacts"]
hint = {c["email"]: c["region_hint"] for c in us_list}
check(hint.get("inph@globex.example") == "India" and hint.get("us1@globex.example") == "", f"looks-like-India hint: {hint}")
x = client.get("/api/vendors?region=India&format=xlsx")
ws = load_workbook(io.BytesIO(x.data)).active
check(x.status_code == 200 and ws.title == "India Vendors" and sorted(r[2] for r in ws.iter_rows(min_row=2, values_only=True)) ==
      ["both@globex.example", "in1@globex.example"], "Excel download of the India list")

r = client.put(f"/api/vendors/{ID[('both@globex.example', 'USA')]}", json={"region": "India"})
check(r.status_code == 400 and "already in the India list" in r.get_json()["error"], f"move into a list that has it: {r.get_json()}")
r = client.put(f"/api/vendors/{ID[('inph@globex.example', 'USA')]}", json={"region": "India"})
check(r.status_code == 200 and r.get_json()["contact"]["region"] == "India", "move a contact to India")
client.put(f"/api/vendors/{ID[('inph@globex.example', 'USA')]}", json={"region": "USA"})

r = client.post("/api/vendors", json={"email": "new@globex.example", "company": "Globex Test", "candidate_id": in_cid})
check(r.status_code == 200 and r.get_json()["region"] == "India", f"saved from an India consultant's window -> India list: {r.get_json()}")
r = client.post("/api/vendors", json={"email": "new@globex.example", "company": "Globex Test", "candidate_id": in_cid})
check(r.status_code == 400 and "India list" in r.get_json()["error"], f"duplicate in India: {r.get_json()}")
r = client.post("/api/vendors", json={"email": "new@globex.example", "company": "Globex Test", "region": "USA"})
check(r.status_code == 200 and r.get_json()["region"] == "USA", "same email can go in the USA list")

# ---- 4. drafts use the consultant's list only
m = client.post("/api/vendors/match", json={"emails": ["poster@globex.example"], "candidate_id": in_cid}).get_json()
check(m["region"] == "India" and sorted(c["email"] for c in m["contacts"]) == ["both@globex.example", "in1@globex.example", "new@globex.example"],
      f"India consultant -> India vendors: {m['contacts']}")
m = client.post("/api/vendors/match", json={"emails": ["poster@globex.example"], "candidate_id": us_cid}).get_json()
check(m["region"] == "USA" and "in1@globex.example" not in [c["email"] for c in m["contacts"]], "US consultant -> USA vendors")

REQ = "Job Title: Java Developer\nLocation: Hyderabad\nPlease send resumes to poster@globex.example\n"
ticked = [ID[("us1@globex.example", "USA")], ID[("both@globex.example", "USA")], ID[("in1@globex.example", "India")], ID[("both@globex.example", "India")]]
r = client.post("/api/outreach/paste-and-draft", json={"candidate_id": in_cid, "raw_jd_text": REQ, "recruiter_email": "poster@globex.example",
                                                         "job_title": "Java Developer", "bcc_contact_ids": ticked})
check(r.status_code == 200 and r.get_json().get("success"), f"India draft: {r.get_json()}")
bcc = sorted(a.strip() for a in (email.message_from_bytes(captured[-1])["Bcc"] or "").split(",") if a.strip())
check(bcc == ["both@globex.example", "in1@globex.example"], f"India draft BCCs only India vendors, even if US ones are sent: {bcc}")
check(r.get_json().get("vendor", {}).get("status") == "created" and r.get_json()["vendor"]["region"] == "India",
      f"the requirement's recruiter is saved to the India list: {r.get_json().get('vendor')}")

jobs = [{"title": "A", "company": "Globex Test", "recruiter_email": "x@globex.example", "country": "India"},
        {"title": "B", "company": "Globex Test", "recruiter_email": "x@globex.example", "country": "United States"},
        {"title": "C", "company": "Nobody Test", "recruiter_email": "x@nobody.example", "country": "India"}]
app_module._tag_known_vendors(jobs, ADMIN)
check(jobs[0].get("vendor", {}).get("count") == 4 and jobs[1].get("vendor", {}).get("count") == 4 and "vendor" not in jobs[2],
      f"job tag counts its own region's list: {[j.get('vendor') for j in jobs]}")

# Postgres rows (models.DbRow) iterate their column names - the counts must read rows by position
class _PgLikeCursor:
    def execute(self, sql, params=()):
        pass

    def fetchall(self):
        return [models.DbRow(["region", "count"], ("USA", 3)), models.DbRow(["region", "count"], ("India", 2))]


class _PgLikeConn:
    def cursor(self):
        return _PgLikeCursor()


check(vendors.region_counts(_PgLikeConn(), ADMIN) == {"USA": 3, "India": 2}, "counts work with Postgres-style rows")

with client.session_transaction() as s:
    s["user"] = US_REC
check(client.get("/api/vendors?region=India").get_json()["contacts"] == [], "a recruiter still sees only their own contacts")
check(client.get("/api/vendors").get_json()["region"] == "USA", "a USA recruiter opens on the USA list")

# ---- 6. the screen
with socket.socket() as sk:
    sk.bind(("127.0.0.1", 0))
    port = sk.getsockname()[1]
server = make_server("127.0.0.1", port, app_module.app, threaded=True)
threading.Thread(target=server.serve_forever, daemon=True).start()
js_errors, shot = [], os.path.join(tmp_dir, "vendor_regions.png")
try:
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="msedge", headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 950})
        page.on("pageerror", lambda e: js_errors.append(str(e)))
        page.goto(f"http://127.0.0.1:{port}/login")
        page.fill("input[name=email]", EMAIL)
        page.fill("input[name=password]", PASSWORD)
        page.click("button[type=submit]")
        page.wait_for_url("**/dashboard**")
        check("app.js?v=5.99.5" in page.content(), "cache-buster not bumped to 5.99.5")
        page.click("a.nav-item[data-tab=vendors]")
        page.wait_for_function("document.getElementById('vendors-count').innerText.includes('USA list')", timeout=10000)
        check(page.inner_text("#vendors-tab-usa") == "USA vendors (4)" and page.inner_text("#vendors-tab-india") == "India vendors (4)",
              f"tab counts: {page.inner_text('#vendors-tab-usa')!r} / {page.inner_text('#vendors-tab-india')!r}")
        check("Move to India" in page.inner_text("#vendors-body"), "hint shown on the Indian-phone contact")
        page.click("#vendors-tab-india")
        page.wait_for_function("document.getElementById('vendors-count').innerText.startsWith('India list: 4')", timeout=10000)
        check("in1@globex.example" in page.inner_text("#vendors-body") and "us1@globex.example" not in page.inner_text("#vendors-body"),
              "India tab shows the India list only")
        check("region=India" in page.get_attribute("#vendors-export", "href"), "Excel link follows the open list")
        page.click("#vendors-add-btn")
        check(page.input_value("#vf-region") == "India", "Add contact defaults to the open list")
        page.fill("#vf-company", "Initech Test")
        page.fill("#vf-email", "ui-add@initech.example")
        page.click("#vf-save")
        page.wait_for_function("document.getElementById('vendors-count').innerText.startsWith('India list: 5')", timeout=10000)
        page.click("#vendors-tab-usa")
        page.wait_for_function("document.getElementById('vendors-count').innerText.startsWith('USA list')", timeout=10000)
        page.click("#vendors-body .vendor-region-hint a")
        page.wait_for_function("document.getElementById('vendors-tab-india').innerText === 'India vendors (6)'", timeout=10000)
        page.screenshot(path=shot)
        browser.close()
finally:
    server.shutdown()

check(not js_errors, f"JS errors: {js_errors}")
if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    print("screenshot:", shot)
    sys.exit(1)
print(f"PASS: vendor lists by region - guessing, one-time sort, separate lists, drafts/BCC/job tags use the consultant's list, Excel per list, Edge tabs. Screenshot: {shot}")
