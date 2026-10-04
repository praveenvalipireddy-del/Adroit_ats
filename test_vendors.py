"""Vendors database: Excel/CSV upload (header auto-detection, preview, duplicates, problems), privacy
(a recruiter sees / edits only their own contacts, admins see all), the import re-check on the
server, and matching a company by email domain or name (own active contacts only).
Standalone; throwaway SQLite DB; no network. People / companies are labelled test fixtures.
Run: python test_vendors.py
"""
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_vendors.db")
os.environ["DATABASE_URL"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import models  # noqa: E402
import vendors  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


def uid(u):
    return u["id"] if isinstance(u, dict) else u


ADMIN = {"id": uid(models.create_user("Vendor Admin", "v-admin@example.invalid", "Vend-1", role="Admin")), "name": "Vendor Admin", "role": "Admin"}
ASHA = {"id": uid(models.create_user("Asha Vendtest", "v-asha@example.invalid", "Vend-2", role="Recruiter")), "name": "Asha Vendtest", "role": "Recruiter"}
RAVI = {"id": uid(models.create_user("Ravi Vendtest", "v-ravi@example.invalid", "Vend-3", role="Recruiter")), "name": "Ravi Vendtest", "role": "Recruiter"}
client = app_module.app.test_client()


def login(u):
    with client.session_transaction() as s:
        s["user"] = u


# ---- helpers
check(vendors.normalize_company("Apex Systems, Inc.") == "apex systems", "normalize Inc")
check(vendors.normalize_company("TEKsystems LLC") == vendors.normalize_company("teksystems"), "normalize LLC")
check(vendors.normalize_company("Smith & Co Pvt Ltd") == "smith and", "normalize & / Pvt Ltd")
check(vendors.company_domain("x@gmail.com") == "" and vendors.company_domain("X@Vendor-One.example") == "vendor-one.example", "company domain")
check(vendors.find_emails("Mail ravi@a-corp.example, cc: jo@a-corp.example. bad@ x@y") == ["ravi@a-corp.example", "jo@a-corp.example"], "find_emails")

# ---- not logged in
check(client.get("/api/vendors").status_code == 401, "list needs login")
check(client.get("/api/vendors/template").status_code == 401, "template needs login")

# ---- template
login(ASHA)
r = client.get("/api/vendors/template")
check(r.status_code == 200 and r.data[:2] == b"PK", "template is an xlsx")
hdr = [c.value for c in load_workbook(io.BytesIO(r.data)).active[1]]
check(hdr == ["Company", "Contact Name", "Email", "Phone", "Title", "Notes"], f"template headers {hdr}")

# ---- Excel upload with unusual headers, a title row above the header, first/last name columns
wb = Workbook()
ws = wb.active
ws.append(["Asha's vendor list (test fixture)"])
ws.append(["Vendor", "First Name", "Last Name", "Mail ID", "Mobile", "Designation"])
ws.append(["Vendor One Test Inc", "Kiran", "Testone", "Kiran@Vendor-One.example", "+1 555 0100", "Recruiter"])
ws.append(["Vendor One Test", "Lata", "Testone", "lata@vendor-one.example", "", ""])
ws.append(["", "Mohan", "Testtwo", "mohan@vendor-two.example", "", ""])          # company from domain
ws.append(["Vendor One Test", "Kiran", "Again", "kiran@vendor-one.example", "", ""])  # repeated in file
ws.append(["Gmail Vendor", "Nina", "Testthree", "nina.vendor@gmail.com", "", ""])   # personal email + company: ok
ws.append(["", "Om", "Testfour", "om.vendor@gmail.com", "", ""])                  # personal, no company: problem
ws.append(["Bad Row Co", "Pia", "Testfive", "not-an-email", "", ""])               # invalid
ws.append([None, None, None, None, None, None])                                    # blank: skipped
buf = io.BytesIO()
wb.save(buf)
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(buf.getvalue()), "asha_vendors.xlsx")},
                content_type="multipart/form-data")
d = r.get_json()
check(r.status_code == 200, f"preview ok: {d}")
check(d["counts"] == {"new": 4, "duplicate": 1, "problem": 2}, f"preview counts {d['counts']}")
by_email = {}
for row in d["rows"]:
    by_email.setdefault(row["email"], row)   # first occurrence (Kiran is repeated on purpose)
check(by_email["kiran@vendor-one.example"]["name"] == "Kiran Testone" and by_email["kiran@vendor-one.example"]["phone"] == "+1 555 0100",
      "first+last name combined, email lower-cased, phone read")
check(by_email["mohan@vendor-two.example"]["company"] == "vendor-two.example", "company from the email domain")
check(by_email["om.vendor@gmail.com"]["status"] == "problem", "personal email with no company is a problem")
check(by_email["not-an-email"]["status"] == "problem", "invalid email is a problem")
check(by_email["nina.vendor@gmail.com"]["status"] == "new", "personal email with a company name is fine")
check(by_email["kiran@vendor-one.example"]["line"] == 3, "row numbers match Excel")
check(models.get_db_connection().cursor().execute("SELECT COUNT(*) FROM vendor_contacts").fetchone()[0] == 0, "preview saves nothing")

# the browser sends back only the new rows; the server re-checks (a forged bad row is not saved)
rows = [row for row in d["rows"] if row["status"] == "new"] + [{"company": "", "email": "forged@gmail.com"}]
r = client.post("/api/vendors/import", json={"rows": rows, "filename": "asha_vendors.xlsx"})
check(r.get_json() == {"success": True, "added": 4, "duplicates": 0, "problems": 1}, f"import result {r.get_json()}")

# uploading the same file again: everything is already in the list
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(buf.getvalue()), "again.xlsx")}, content_type="multipart/form-data")
check(r.get_json()["counts"]["new"] == 0 and r.get_json()["counts"]["duplicate"] == 5, f"re-upload: {r.get_json()['counts']}")

# CSV with BOM; bad file types
csv_bytes = "﻿Company Name,Recruiter Name,Email Address,Phone Number\nVendor Three Test LLC,Sita Testsix,sita@vendor-three.example,9876543210\n".encode("utf-8")
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(csv_bytes), "v.csv")}, content_type="multipart/form-data")
check(r.status_code == 200 and r.get_json()["counts"]["new"] == 1 and r.get_json()["rows"][0]["company"] == "Vendor Three Test LLC", f"csv: {r.get_json()}")
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(b"x"), "v.pdf")}, content_type="multipart/form-data")
check(r.status_code == 400 and "Excel" in r.get_json()["error"], "pdf rejected")
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(b"Name,Phone\nA,1\n"), "v.csv")}, content_type="multipart/form-data")
check(r.status_code == 400 and "No email" in r.get_json()["error"], "no emails at all -> clear error")

# ---- a free-form sheet with NO header row (real-world shape: lists pasted side by side, notes,
# LinkedIn links, a second block with its own headers further down). Regression: the upload used to
# reject it ("no email column"), and a cell like "x@gmail.com" was taken for an "Email" header.
wb = Workbook()
ws = wb.active
ws.append([None])
ws.append(["ana@scan-one.example", None])
ws.append(["ben@scan-one.example", "(732) 555-0101", "SAP Recruiter"])
ws.append([None, None, None, None, None, "note", "cara@scan-two.example", "dev@gmail.com"])
ws.append(["eli@scan-three.example", "https://www.linkedin.com/in/eli-test", "Phone: +1-602-555-0102"])
ws.append([None, "FF", "Recruiter Name", "Mail id", "Contact", "Notes", "LinkedIn"])
ws.append([None, 1, "Fay Test", "fay@scan-four.example", "Works for AT&T clients", "408-555-0103 / 510-555-0104", "linkedin.com/in/fay"])
ws.append([None, 2, None, "ana@scan-one.example", None, None])       # repeated
buf2 = io.BytesIO()
wb.save(buf2)
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(buf2.getvalue()), "Book1.xlsx")}, content_type="multipart/form-data")
d = r.get_json()
check(r.status_code == 200 and d["counts"] == {"new": 5, "duplicate": 1, "problem": 1}, f"free-form sheet counts: {d.get('counts', d)}")
got = {}
for row in d.get("rows", []):
    got.setdefault(row["email"], row)
check(got["ben@scan-one.example"]["phone"] == "(732) 555-0101" and got["ben@scan-one.example"]["notes"] == "SAP Recruiter", f"phone + note beside the email: {got.get('ben@scan-one.example')}")
check(got["ben@scan-one.example"]["company"] == "scan-one.example", "company from the email domain")
check(got["cara@scan-two.example"]["phone"] == "" and got["dev@gmail.com"]["status"] == "problem", "emails in other columns found; personal email without company flagged")
check(got["eli@scan-three.example"]["phone"] == "Phone: +1-602-555-0102" and got["eli@scan-three.example"]["notes"] == "", "LinkedIn link is not a note")
check(got["fay@scan-four.example"]["phone"] == "408-555-0103 / 510-555-0104" and got["fay@scan-four.example"]["notes"] == "Works for AT&T clients", f"second block: {got.get('fay@scan-four.example')}")
r = client.post("/api/vendors/upload-preview", data={"file": (io.BytesIO(b"just some text\nno addresses here\n"), "x.csv")}, content_type="multipart/form-data")
check(r.status_code == 400 and "No email" in r.get_json()["error"], "a file without any email -> clear error")

# ---- manual add (Ravi) - same email as Asha's is fine: lists are per recruiter
login(RAVI)
r = client.post("/api/vendors", json={"company": "Vendor One Test", "name": "Kiran Testone", "email": "kiran@vendor-one.example"})
check(r.status_code == 200 and r.get_json()["id"], f"Ravi adds the same contact to his own list: {r.get_json()}")
r = client.post("/api/vendors", json={"company": "Vendor One Test", "email": "kiran@vendor-one.example"})
check(r.status_code == 400 and "Already" in r.get_json()["error"], "duplicate within own list refused")
r = client.post("/api/vendors", json={"company": "", "email": "x@gmail.com"})
check(r.status_code == 400, "personal email without company refused")

# ---- privacy
asha_list = None
login(ASHA)
asha_list = client.get("/api/vendors").get_json()
check(len(asha_list["contacts"]) == 4 and not asha_list["is_admin"] and asha_list["owners"] == [], f"Asha sees her 4: {len(asha_list['contacts'])}")
check(asha_list["companies"] == 3, f"3 companies: {asha_list['companies']}")
login(RAVI)
ravi_list = client.get("/api/vendors").get_json()["contacts"]
check(len(ravi_list) == 1 and ravi_list[0]["owner_user_id"] == RAVI["id"], "Ravi sees only his own")
check(client.get(f"/api/vendors?owner={ASHA['id']}").get_json()["contacts"][0]["owner_user_id"] == RAVI["id"], "?owner ignored for recruiters")
asha_contact = asha_list["contacts"][0]["id"]
check(client.put(f"/api/vendors/{asha_contact}", json={"phone": "1"}).status_code == 404, "Ravi can't edit Asha's contact")
check(client.delete(f"/api/vendors/{asha_contact}").status_code == 404, "Ravi can't delete Asha's contact")
login(ADMIN)
adm = client.get("/api/vendors").get_json()
check(adm["is_admin"] and len(adm["contacts"]) == 5, f"admin sees all 5: {len(adm['contacts'])}")
check({o["name"] for o in adm["owners"]} == {"Asha Vendtest", "Ravi Vendtest"}, f"owners {adm['owners']}")
check(len(client.get(f"/api/vendors?owner={RAVI['id']}").get_json()["contacts"]) == 1, "admin filters by recruiter")
check(len(client.get("/api/vendors?q=vendor-two").get_json()["contacts"]) == 1, "search by email")

# ---- edit / status / delete (own)
login(ASHA)
lata = next(c for c in asha_list["contacts"] if c["email"] == "lata@vendor-one.example")
r = client.put(f"/api/vendors/{lata['id']}", json={"phone": "+1 555 0101", "title": "Lead Recruiter"})
check(r.status_code == 200 and r.get_json()["contact"]["phone"] == "+1 555 0101", "edit own contact")
check(client.put(f"/api/vendors/{lata['id']}", json={"status": "spam"}).status_code == 400, "unknown status refused")
check(client.put(f"/api/vendors/{lata['id']}", json={"company_name": ""}).status_code == 400, "empty company refused")

# ---- match: domain first, then company name; own active contacts only
r = client.post("/api/vendors/match", json={"company": "", "emails": ["recruiter.x@Vendor-One.example"]}).get_json()
check(r["matched_by"] == "domain" and sorted(c["email"] for c in r["contacts"]) == ["kiran@vendor-one.example", "lata@vendor-one.example"],
      f"domain match: {r}")
check(r["company"] == "Vendor One Test Inc", f"company label: {r['company']}")
client.put(f"/api/vendors/{lata['id']}", json={"status": "unsubscribed"})
r = client.post("/api/vendors/match", json={"emails": "Please reply to recruiter.x@vendor-one.example"}).get_json()
check([c["email"] for c in r["contacts"]] == ["kiran@vendor-one.example"], f"unsubscribed skipped; emails parsed from text: {r}")
r = client.post("/api/vendors/match", json={"company": "VENDOR ONE TEST, LLC", "emails": ["someone@gmail.com"]}).get_json()
check(r["matched_by"] == "company" and [c["email"] for c in r["contacts"]] == ["kiran@vendor-one.example"], f"company-name match (gmail ignored): {r}")
r = client.post("/api/vendors/match", json={"company": "Gmail Vendor", "emails": []}).get_json()
check([c["email"] for c in r["contacts"]] == ["nina.vendor@gmail.com"], "personal-email contact matched by company name")
r = client.post("/api/vendors/match", json={"company": "Unknown Co", "emails": ["a@unknown.example"]}).get_json()
check(r["contacts"] == [] and r["matched_by"] in ("", "company", "domain"), "no match")
login(RAVI)
r = client.post("/api/vendors/match", json={"emails": ["x@vendor-two.example"]}).get_json()
check(r["contacts"] == [], "Ravi never matches Asha's contacts")
login(ADMIN)
r = client.post("/api/vendors/match", json={"emails": ["x@vendor-two.example"]}).get_json()
check(r["contacts"] == [], "even an admin's drafts use only the admin's own contacts")

# ---- mark_emailed + delete
conn = models.get_db_connection()
vendors.mark_emailed(conn, [lata["id"]], "Java Developer - test")
conn.close()
login(ASHA)
c = next(c for c in client.get("/api/vendors").get_json()["contacts"] if c["id"] == lata["id"])
check(c["last_emailed_at"] and c["last_emailed_note"] == "Java Developer - test", "last emailed recorded")
check(client.delete(f"/api/vendors/{lata['id']}").status_code == 200 and len(client.get("/api/vendors").get_json()["contacts"]) == 3, "delete own")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: vendors - Excel/CSV upload + preview + server re-check, private per recruiter, admins see all, domain/company match.")
