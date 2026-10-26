"""Vendor contacts: each recruiter's own list of vendor recruiters (company, name, email, phone...).

- Private: a recruiter sees and uses only the contacts they added; admins can see everyone's (for
  oversight) but a draft only ever uses the drafting recruiter's own contacts.
- Excel / CSV upload with header auto-detection ("Mail ID", "Vendor", "Mobile"...), a preview
  (new / already in your list / problem) and an import of the good rows only.
- A company is recognised by its email domain (most reliable - e.g. @apexsystems.com) or by its
  name with Inc / LLC / Corp ... removed. Personal domains (gmail, yahoo...) never identify a company.
- Contacts can be marked unsubscribed / bounced; those are never put in an email.
- Two lists per recruiter: USA vendors and India vendors (vendor_contacts.region). A draft only ever
  uses the list of its consultant's region, so an India requirement never BCCs a US vendor. The same
  email may sit in both lists. Existing contacts were sorted once by guess_region() (phone / email),
  falling back to the recruiter's own team.
"""
import csv
import io
import re
from typing import Dict, Iterable, List, Optional, Tuple

PERSONAL_DOMAINS = {
    "gmail.com", "googlemail.com", "yahoo.com", "yahoo.in", "yahoo.co.in", "ymail.com", "hotmail.com",
    "outlook.com", "live.com", "msn.com", "aol.com", "icloud.com", "me.com", "mac.com", "protonmail.com",
    "proton.me", "zoho.com", "zohomail.com", "zohomail.in", "rediffmail.com", "gmx.com", "mail.com", "yandex.com",
}
STATUSES = ["active", "unsubscribed", "bounced"]
MAX_ROWS = 5000
MAX_BCC = 10   # vendor contacts BCC'd on one draft
_EMAIL_RE = re.compile(r"^[^@\s,;<>]+@[^@\s,;<>]+\.[A-Za-z]{2,}$")
_EMAIL_FIND_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_SUFFIXES = {"inc", "llc", "ltd", "limited", "corp", "corporation", "co", "company", "pvt", "private", "plc",
             "llp", "lp", "gmbh", "incorporated", "the"}

HEADER_SYNONYMS = {
    "company": ["company", "company name", "vendor", "vendor name", "vendor company", "organization", "organisation",
                "firm", "employer", "client", "account"],
    "name": ["name", "contact", "contact name", "recruiter", "recruiter name", "full name", "person", "poc"],
    "first_name": ["first name", "firstname", "first"],
    "last_name": ["last name", "lastname", "last", "surname"],
    "email": ["email", "e-mail", "e mail", "mail", "mail id", "email id", "email address", "emailid", "mailid",
              "recruiter email", "contact email"],
    "phone": ["phone", "mobile", "cell", "contact number", "phone number", "mobile number", "tel", "telephone",
              "contact no", "phone no", "mobile no", "whatsapp"],
    "title": ["title", "designation", "role", "position", "job title"],
    "notes": ["notes", "note", "comments", "comment", "remarks"],
}


class VendorError(ValueError):
    pass


REGIONS = ("USA", "India")


def norm_region(value: Optional[str], default: str = "USA") -> str:
    v = (value or "").strip().lower()
    return "India" if v in ("india", "in") else ("USA" if v in ("usa", "us", "united states") else default)


_US_GROUPED = re.compile(r"^(\(\d{3}\)[\s.-]*|\d{3}[\s.-]+)\d{3}[\s.-]+\d{4}$")
_IN_GROUPED = re.compile(r"^\d{5}[\s.-]+\d{5}$")


def guess_region(phone: Optional[str], email: Optional[str]) -> str:
    """'India' / 'USA' when the phone or email says so, '' when they don't.
    +91 / 91XXXXXXXXXX / '98765 43210' / an .in email -> India; +1 / 1XXXXXXXXXX / '(949) 688-0755' -> USA.
    A bare 10-digit number is NOT a signal: US area codes can start with 6-9 too (949...)."""
    p = (phone or "").strip()
    digits = re.sub(r"\D", "", p)
    if p.startswith("+91") or p.startswith("0091") or (len(digits) == 12 and digits.startswith("91") and digits[2] in "6789"):
        return "India"
    if p.startswith("+1") or (len(digits) == 11 and digits.startswith("1")):
        return "USA"
    if _IN_GROUPED.match(p):
        return "India"
    if _US_GROUPED.match(p):
        return "USA"
    d = email_domain(email)
    if d.endswith(".in"):
        return "India"
    return ""


def backfill_regions(conn) -> int:
    """Give every contact without a region one, once: from its phone / email, else its recruiter's team."""
    cur = conn.cursor()
    cur.execute("""SELECT v.id, v.phone, v.email, u.region FROM vendor_contacts v LEFT JOIN users u ON u.id = v.owner_user_id
                   WHERE v.region IS NULL OR v.region = ''""")
    rows = cur.fetchall()
    for r in rows:
        region = guess_region(r[1], r[2]) or norm_region(r[3])
        cur.execute("UPDATE vendor_contacts SET region = ? WHERE id = ?", (region, r[0]))
    conn.commit()
    return len(rows)


def normalize_company(name: Optional[str]) -> str:
    """"Apex Systems, Inc." -> "apex systems"."""
    s = re.sub(r"[^a-z0-9& ]+", " ", (name or "").lower().replace("&", " and "))
    words = [w for w in s.split() if w not in _SUFFIXES]
    return " ".join(words)


def email_domain(email: Optional[str]) -> str:
    e = (email or "").strip().lower()
    return e.rsplit("@", 1)[1] if "@" in e else ""


def company_domain(email: Optional[str]) -> str:
    """The email's domain if it identifies a company ('' for gmail/yahoo/...)."""
    d = email_domain(email)
    return "" if (not d or d in PERSONAL_DOMAINS) else d


def find_emails(text: str) -> List[str]:
    seen, out = set(), []
    for m in _EMAIL_FIND_RE.findall(text or ""):
        e = m.strip(".'").lower()
        if e not in seen and _EMAIL_RE.match(e):
            seen.add(e)
            out.append(e)
    return out


# ---------------------------------------------------------------- upload parsing

def _norm_header(h) -> str:
    return re.sub(r"[^a-z ]+", " ", str(h or "").lower()).strip()


def _map_headers(headers: List) -> Dict[str, int]:
    norm = [re.sub(r"\s+", " ", _norm_header(h)) for h in headers]
    mapping = {}
    for field, names in HEADER_SYNONYMS.items():
        for i, h in enumerate(norm):
            if i in mapping.values():
                continue
            if h in names:
                mapping[field] = i
                break
    if "email" not in mapping:   # looser: any header containing "mail"
        for i, h in enumerate(norm):
            if "mail" in h and i not in mapping.values():
                mapping["email"] = i
                break
    return mapping


def read_table(file_bytes: bytes, filename: str) -> List[List]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        except Exception:
            raise VendorError("That file couldn't be opened as an Excel workbook.")
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    raise VendorError("Upload an Excel (.xlsx) or CSV file.")


def parse_rows(table: List[List]) -> List[Dict]:
    """First row with an email-like header is the header row; returns raw row dicts (unvalidated)."""
    header_idx, mapping = None, {}
    for i, row in enumerate(table[:20]):
        if any(find_emails(_cell_text(v)) for v in (row or [])):
            continue   # a row holding real addresses ("...@gmail.com" contains "mail") is data, not headers
        m = _map_headers(row or [])
        if "email" in m:
            header_idx, mapping = i, m
            break
    if header_idx is None:
        return scan_rows(table)
    # Trust the header row only when its email column holds (nearly) all of the sheet's emails;
    # otherwise it's a small block inside a free-form sheet and the other emails would be lost.
    col = mapping["email"]
    total = sum(len(find_emails(_cell_text(v))) for row in table for v in (row or []))
    in_col = sum(len(find_emails(_cell_text(row[col]))) for row in table[header_idx + 1:] if row and col < len(row))
    if in_col < 0.9 * total:
        return scan_rows(table)
    out = []
    for n, row in enumerate(table[header_idx + 1:], start=header_idx + 2):
        row = list(row or [])

        def cell(field):
            i = mapping.get(field)
            v = row[i] if i is not None and i < len(row) else None
            return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""

        if not any(str(c or "").strip() for c in row):
            continue
        name = cell("name") or " ".join(x for x in (cell("first_name"), cell("last_name")) if x)
        out.append({"line": n, "company": cell("company"), "name": name, "email": cell("email").lower(),
                    "phone": cell("phone"), "title": cell("title"), "notes": cell("notes")})
        if len(out) > MAX_ROWS:
            raise VendorError(f"That file has more than {MAX_ROWS:,} rows - split it into smaller files.")
    return out


_URL_RE = re.compile(r"^(https?://|www\.|[a-z0-9-]+\.[a-z]{2,}/)", re.I)


def _cell_text(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def _looks_like_phone(text: str) -> bool:
    digits = re.sub(r"\D", "", text)
    return 10 <= len(digits) <= 26 and not _URL_RE.match(text) and "@" not in text \
        and len(re.sub(r"[\d\s()+\-./#*:,;xXextEXTPphoneMmobileDdirectCcellTtel]", "", text)) <= 3


def scan_rows(table: List[List]) -> List[Dict]:
    """A sheet without a header row (lists pasted side by side, notes in between): every email in
    any cell becomes a contact. Phone = the first phone-looking cell to its right (before the next
    email); other text cells there (not links) become notes. Company comes from the email domain
    in validate() - names aren't guessed, since a free-form sheet doesn't say which cell is a name."""
    out = []
    for n, row in enumerate(table, start=1):
        cells = [_cell_text(v) for v in (row or [])]
        for i, text in enumerate(cells):
            found = find_emails(text)
            if not found:
                continue
            phone, notes = "", []
            for nxt in cells[i + 1:i + 3]:
                if not nxt:
                    continue
                if find_emails(nxt):
                    break
                if not phone and _looks_like_phone(nxt):
                    phone = nxt
                elif not _URL_RE.match(nxt) and not re.fullmatch(r"\d{1,3}", nxt):
                    notes.append(nxt)
            for e in found:
                out.append({"line": n, "company": "", "name": "", "email": e, "phone": phone[:60],
                            "title": "", "notes": "; ".join(notes)[:300]})
        if len(out) > MAX_ROWS:
            raise VendorError(f"That file has more than {MAX_ROWS:,} emails - split it into smaller files.")
    if not out:
        raise VendorError("No email addresses found in that file.")
    return out


def validate(rows: Iterable[Dict], existing_emails: set) -> List[Dict]:
    """Each row + 'status' ('new' | 'duplicate' | 'problem') and 'reason'. Company falls back to the
    email's company domain when the Company cell is empty."""
    seen, out = set(), []
    for r in rows:
        r = dict(r)
        email = (r.get("email") or "").strip().lower()
        r["email"] = email
        if not r.get("company") and company_domain(email):
            r["company"] = company_domain(email)
        if not email or not _EMAIL_RE.match(email):
            r["status"], r["reason"] = "problem", "Missing or invalid email"
        elif not r.get("company"):
            r["status"], r["reason"] = "problem", "Company missing (personal email address)"
        elif email in existing_emails:
            r["status"], r["reason"] = "duplicate", "Already in your list"
        elif email in seen:
            r["status"], r["reason"] = "duplicate", "Repeated in this file"
        else:
            r["status"], r["reason"] = "new", ""
            seen.add(email)
        for k in ("company", "name", "phone", "title", "notes"):
            r[k] = (r.get(k) or "")[:300]
        out.append(r)
    return out


# ---------------------------------------------------------------- database

def _owner_emails(cur, owner_id: int, region: str = "USA") -> set:
    cur.execute("SELECT LOWER(email) FROM vendor_contacts WHERE owner_user_id = ? AND region = ?", (owner_id, norm_region(region)))
    return {r[0] for r in cur.fetchall()}


def preview(conn, owner_id: int, file_bytes: bytes, filename: str, region: str = "USA") -> List[Dict]:
    rows = parse_rows(read_table(file_bytes, filename))
    return validate(rows, _owner_emails(conn.cursor(), owner_id, region))


def import_rows(conn, owner_id: int, rows: List[Dict], source: str = "upload", region: str = "USA") -> Dict:
    """Re-validates on the server and inserts only the new rows into that region's list. Returns counts."""
    region = norm_region(region)
    cur = conn.cursor()
    checked = validate(rows, _owner_emails(cur, owner_id, region))
    added = 0
    for r in checked:
        if r["status"] != "new":
            continue
        cur.execute("""INSERT INTO vendor_contacts (owner_user_id, company_name, company_norm, email, email_domain,
                       contact_name, phone, title, notes, status, source, region) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', ?, ?)""",
                    (owner_id, r["company"], normalize_company(r["company"]), r["email"], company_domain(r["email"]),
                     r.get("name") or "", r.get("phone") or "", r.get("title") or "", r.get("notes") or "", source, region))
        added += 1
    conn.commit()
    return {"added": added, "duplicates": sum(1 for r in checked if r["status"] == "duplicate"),
            "problems": sum(1 for r in checked if r["status"] == "problem")}


def list_contacts(conn, user: Dict, q: str = "", owner_id: Optional[int] = None, region: Optional[str] = None) -> List[Dict]:
    """A recruiter gets their own contacts; an admin gets everyone's (optionally one recruiter's).
    region: only that list (USA / India)."""
    is_admin = "Admin" in (user.get("role") or "")
    where, params = [], []
    if region:
        where.append("v.region = ?")
        params.append(norm_region(region))
    if not is_admin:
        where.append("v.owner_user_id = ?")
        params.append(user["id"])
    elif owner_id:
        where.append("v.owner_user_id = ?")
        params.append(owner_id)
    if q:
        like = f"%{q.lower()}%"
        where.append("(LOWER(v.company_name) LIKE ? OR LOWER(v.email) LIKE ? OR LOWER(COALESCE(v.contact_name, '')) LIKE ?)")
        params += [like, like, like]
    cur = conn.cursor()
    cur.execute(f"""SELECT v.*, u.name AS owner_name FROM vendor_contacts v LEFT JOIN users u ON u.id = v.owner_user_id
                    {('WHERE ' + ' AND '.join(where)) if where else ''}
                    ORDER BY LOWER(v.company_name), LOWER(COALESCE(v.contact_name, v.email)) LIMIT 5000""", params)
    out = []
    for r in cur.fetchall():
        d = dict(r)
        hint = guess_region(d.get("phone"), d.get("email"))
        d["region_hint"] = hint if hint and hint != d.get("region") else ""   # "looks like an India contact"
        out.append(d)
    return out


def region_counts(conn, user: Dict, owner_id: Optional[int] = None) -> Dict[str, int]:
    """How many contacts each list holds (what the USA / India tabs show)."""
    is_admin = "Admin" in (user.get("role") or "")
    where, params = [], []
    if not is_admin:
        where.append("owner_user_id = ?")
        params.append(user["id"])
    elif owner_id:
        where.append("owner_user_id = ?")
        params.append(owner_id)
    cur = conn.cursor()
    cur.execute(f"SELECT region, COUNT(*) FROM vendor_contacts {('WHERE ' + ' AND '.join(where)) if where else ''} GROUP BY region", params)
    counts = {r: 0 for r in REGIONS}
    for row in cur.fetchall():          # by position: a Postgres row (models.DbRow) iterates its column NAMES
        reg = norm_region(row[0])
        counts[reg] = counts.get(reg, 0) + int(row[1] or 0)
    return counts


def _contact_for_edit(cur, contact_id: int, user: Dict) -> Dict:
    cur.execute("SELECT * FROM vendor_contacts WHERE id = ?", (contact_id,))
    row = cur.fetchone()
    if not row:
        raise LookupError("Contact not found")
    row = dict(row)
    if row["owner_user_id"] != user["id"] and "Admin" not in (user.get("role") or ""):
        raise LookupError("Contact not found")   # private: other recruiters' contacts don't exist for you
    return row


def update_contact(conn, contact_id: int, user: Dict, data: Dict) -> Dict:
    cur = conn.cursor()
    row = _contact_for_edit(cur, contact_id, user)
    fields = {}
    for k in ("company_name", "contact_name", "phone", "title", "notes"):
        if k in data:
            fields[k] = str(data.get(k) or "").strip()[:300]
    if "company_name" in fields:
        if not fields["company_name"]:
            raise VendorError("Company can't be empty.")
        fields["company_norm"] = normalize_company(fields["company_name"])
        fields["company_guess"] = 0              # the recruiter set the company: no longer a guess
    if "status" in data:
        if data["status"] not in STATUSES:
            raise VendorError("Unknown status.")
        fields["status"] = data["status"]
    if "region" in data:
        new_region = norm_region(data.get("region"), default="")
        if not new_region:
            raise VendorError("Region must be USA or India.")
        if new_region != row.get("region"):
            if row["email"].lower() in _owner_emails(cur, row["owner_user_id"], new_region):
                raise VendorError(f"{row['email']} is already in the {new_region} list.")
            fields["region"] = new_region
    if fields:
        cur.execute(f"UPDATE vendor_contacts SET {', '.join(k + ' = ?' for k in fields)} WHERE id = ?",
                    list(fields.values()) + [contact_id])
        conn.commit()
    row.update(fields)
    return row


def delete_contact(conn, contact_id: int, user: Dict) -> None:
    cur = conn.cursor()
    _contact_for_edit(cur, contact_id, user)
    cur.execute("DELETE FROM vendor_contacts WHERE id = ?", (contact_id,))
    conn.commit()


def add_contact(conn, owner_id: int, data: Dict, source: str = "manual", region: str = "USA") -> Tuple[Optional[int], str]:
    """Add one contact to that region's list; returns (id, '') or (None, reason)."""
    region = norm_region(region)
    row = {"company": str(data.get("company") or "").strip(), "name": str(data.get("name") or "").strip(),
           "email": str(data.get("email") or "").strip().lower(), "phone": str(data.get("phone") or "").strip(),
           "title": str(data.get("title") or "").strip(), "notes": str(data.get("notes") or "").strip()}
    checked = validate([row], _owner_emails(conn.cursor(), owner_id, region))[0]
    if checked["status"] != "new":
        return None, (f"Already in your {region} list" if checked["status"] == "duplicate" else checked["reason"])
    import_rows(conn, owner_id, [checked], source=source, region=region)
    cur = conn.cursor()
    cur.execute("SELECT id FROM vendor_contacts WHERE owner_user_id = ? AND region = ? AND LOWER(email) = ?",
                (owner_id, region, checked["email"]))
    r = cur.fetchone()
    return (r[0] if r else None), ""


def match(conn, owner_id: int, company: str = "", emails: Iterable[str] = (), region: str = "USA") -> Dict:
    """This recruiter's active contacts (in that region's list) at the company of a job / pasted
    requirement. Email domain first (most reliable), then the normalised company name."""
    region = norm_region(region)
    cur = conn.cursor()
    domains = [d for d in {company_domain(e) for e in emails or []} if d]
    rows, how = [], ""
    if domains:
        cur.execute(f"""SELECT * FROM vendor_contacts WHERE owner_user_id = ? AND region = ? AND status = 'active'
                        AND email_domain IN ({", ".join(["?"] * len(domains))}) ORDER BY id""", (owner_id, region, *domains))
        rows, how = [dict(r) for r in cur.fetchall()], "domain"
    norm = normalize_company(company)
    if not rows and norm:
        cur.execute("SELECT * FROM vendor_contacts WHERE owner_user_id = ? AND region = ? AND status = 'active' AND company_norm = ? ORDER BY id",
                    (owner_id, region, norm))
        rows, how = [dict(r) for r in cur.fetchall()], "company"
    label = rows[0]["company_name"] if rows else ""
    return {"company": label, "matched_by": how, "contacts": rows, "region": region}


def unknown_emails(conn, owner_id: int, emails: Iterable[str], region: str = "USA") -> List[str]:
    """Emails (e.g. found in a pasted requirement) that aren't in this recruiter's list of that region."""
    have = _owner_emails(conn.cursor(), owner_id, region)
    return [e for e in dict.fromkeys((x or "").strip().lower() for x in emails or []) if e and e not in have]


def contacts_for_bcc(conn, owner_id: int, contact_ids: Iterable, exclude_emails: Iterable[str] = (),
                     region: str = "USA") -> List[Dict]:
    """The contacts a draft may BCC: only this recruiter's own, only active ones, only from the
    consultant's region list, never the To address, at most MAX_BCC (in the order given)."""
    ids = []
    for i in contact_ids or []:
        try:
            ids.append(int(i))
        except (TypeError, ValueError):
            continue
    ids = list(dict.fromkeys(ids))
    if not ids:
        return []
    cur = conn.cursor()
    cur.execute(f"""SELECT * FROM vendor_contacts WHERE owner_user_id = ? AND region = ? AND status = 'active'
                    AND id IN ({", ".join(["?"] * len(ids))})""", (owner_id, norm_region(region), *ids))
    by_id = {r["id"]: dict(r) for r in cur.fetchall()}
    skip = {(e or "").strip().lower() for e in exclude_emails or []}
    out, seen = [], set()
    for i in ids:
        c = by_id.get(i)
        if c and c["email"].lower() not in skip and c["email"].lower() not in seen:
            seen.add(c["email"].lower())
            out.append(c)
    return out[:MAX_BCC]


def vendor_index(conn, owner_id: int, region: str = "USA") -> Dict:
    """This recruiter's active contacts of one region grouped by email domain and by normalised
    company - built once, then match_job() checks every job of a search without another query."""
    cur = conn.cursor()
    cur.execute("SELECT * FROM vendor_contacts WHERE owner_user_id = ? AND region = ? AND status = 'active' ORDER BY id",
                (owner_id, norm_region(region)))
    idx = {"domains": {}, "companies": {}}
    for r in cur.fetchall():
        r = dict(r)
        if r.get("email_domain"):
            idx["domains"].setdefault(r["email_domain"], []).append(r)
        if r.get("company_norm"):
            idx["companies"].setdefault(r["company_norm"], []).append(r)
    return idx


def match_job(idx: Dict, job: Dict) -> Tuple[str, List[Dict]]:
    """(vendor company, contacts) for a job: recruiter email domain first, then company name."""
    d = company_domain(job.get("recruiter_email"))
    rows = idx["domains"].get(d, []) if d else []
    if not rows:
        rows = idx["companies"].get(normalize_company(job.get("company")), []) if normalize_company(job.get("company")) else []
    return (rows[0]["company_name"] if rows else ""), rows


def bcc_for_job(idx: Dict, job: Dict, to_email: str) -> List[Dict]:
    """The contacts to BCC on a draft for this job: matched, never the To address, at most MAX_BCC."""
    _, rows = match_job(idx, job)
    to = (to_email or "").strip().lower()
    out, seen = [], set()
    for r in rows:
        e = r["email"].lower()
        if e != to and e not in seen:
            seen.add(e)
            out.append(r)
    return out[:MAX_BCC]


def mark_emailed(conn, contact_ids: Iterable[int], note: str) -> None:
    ids = [int(i) for i in contact_ids or []]
    if not ids:
        return
    conn.cursor().execute(f"""UPDATE vendor_contacts SET last_emailed_at = CURRENT_TIMESTAMP, last_emailed_note = ?
                              WHERE id IN ({", ".join(["?"] * len(ids))})""", (note[:300], *ids))
    conn.commit()


def export_xlsx(rows: List[Dict], region: str) -> bytes:
    """One region's list as an Excel file (what the Vendors tab shows)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = f"{norm_region(region)} Vendors"
    ws.append(["Company", "Contact Name", "Email", "Phone", "Title", "Notes", "Status", "Last emailed"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for r in rows:
        ws.append([r.get("company_name") or "", r.get("contact_name") or "", r.get("email") or "", r.get("phone") or "",
                   r.get("title") or "", r.get("notes") or "", r.get("status") or "", str(r.get("last_emailed_at") or "")[:10]])
    for col, w in zip("ABCDEFGH", (28, 24, 34, 18, 22, 40, 14, 14)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"
    ws.append(["Company", "Contact Name", "Email", "Phone", "Title", "Notes"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for col, w in zip("ABCDEF", (28, 24, 34, 18, 22, 40)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- automatic capture from job requirements
# A vendor recruiter's email entered on a job, used for a draft, or read from a Gmail requirement is saved to the
# recruiter's own vendor list automatically - with the company name and everything the posting tells us. Job-board
# relay addresses (Dice private email, Indeed, Naukri ...) and personal addresses are never saved.
BOARD_DOMAINS = ("dice.com", "indeed.com", "indeedemail.com", "naukri.com", "naukrigulf.com", "linkedin.com", "foundit.in",
                 "monster.com", "ziprecruiter.com", "glassdoor.com", "shine.com", "timesjobs.com", "apna.co", "internshala.com")
_PLACEHOLDER_COMPANIES = {"", "employer", "(see post)", "direct client", "confidential", "unknown", "n/a", "na"}


def _is_board(email: str) -> bool:
    d = email_domain(email)
    return any(d == b or d.endswith("." + b) for b in BOARD_DOMAINS)


def _domain_label(email: str) -> str:
    """'priya@pyramidci.com' -> 'pyramidci' (the part that names the company)."""
    d = company_domain(email)
    parts = [p for p in d.split(".") if p]
    if len(parts) >= 3 and parts[-2] in ("co", "com", "net", "org", "ac") and len(parts[-1]) == 2:
        return parts[-3]               # tcs.co.in -> tcs
    return parts[-2] if len(parts) >= 2 else (parts[0] if parts else "")


def company_matches_email(company: str, email: str) -> bool:
    """Does this company name belong to this email's domain? ('Pyramid Consulting' / pyramidci.com -> yes,
    'Bank of America' / pyramidci.com -> no)."""
    label = re.sub(r"[^a-z0-9]", "", _domain_label(email))
    words = [w for w in normalize_company(company).split() if w not in ("and", "of", "technologies", "technology", "solutions",
                                                                          "consulting", "services", "systems", "group", "global",
                                                                          "it", "software", "staffing", "infotech")]
    compact = re.sub(r"[^a-z0-9]", "", normalize_company(company))
    if not label or not compact:
        return False
    if label in compact or compact in label:
        return True
    initials = "".join(w[0] for w in normalize_company(company).split() if w)
    if len(initials) >= 2 and label.startswith(initials):
        return True                   # 'Tata Consultancy Services' / tcs.com
    return any(len(w) >= 4 and (w in label or label.startswith(w[:5])) for w in words)


def pretty_domain(email: str) -> str:
    label = _domain_label(email)
    return label[:1].upper() + label[1:] if label else ""


def resolve_company(email: str, vendor_company: str = "", posting_company: str = "", client: str = "") -> Tuple[str, bool, str]:
    """(company name, is a guess, client). The vendor company the AI read wins; then the posting's company when it
    belongs to the email's domain (otherwise that company is the end client); else a readable domain, as a guess."""
    vendor_company = (vendor_company or "").strip()
    posting = (posting_company or "").strip()
    client = (client or "").strip()
    if posting.lower() in _PLACEHOLDER_COMPANIES:
        posting = ""
    if vendor_company and vendor_company.lower() not in _PLACEHOLDER_COMPANIES:
        if posting and not client and normalize_company(posting) != normalize_company(vendor_company):
            client = posting
        return vendor_company[:200], False, client
    if posting and company_matches_email(posting, email):
        return posting[:200], False, client
    if posting and not client:
        client = posting
    return pretty_domain(email), True, client


def _find(cur, owner_id: int, region: str, email: str):
    cur.execute("SELECT * FROM vendor_contacts WHERE owner_user_id = ? AND region = ? AND LOWER(email) = ?", (owner_id, region, email))
    r = cur.fetchone()
    return dict(r) if r else None


def capture(conn, owner_id: int, region: str, email: str, *, vendor_company: str = "", posting_company: str = "", client: str = "",
            name: str = "", phone: str = "", title: str = "", linkedin_url: str = "", job_title: str = "", job_url: str = "",
            found_via: str = "") -> Dict:
    """Save / enrich one vendor contact from a requirement. Returns {status: created | updated | unchanged | skipped,
    id, company, company_guess, name, region, reason}. Never overwrites what the recruiter typed: only empty fields
    are filled, and a guessed company name is replaced by a real one."""
    region = norm_region(region)
    email = (email or "").strip().lower()
    out = {"status": "skipped", "id": None, "company": "", "company_guess": False, "name": name or "", "region": region, "reason": ""}
    if not owner_id or not email or not _EMAIL_RE.match(email):
        out["reason"] = "no valid email"
        return out
    if not company_domain(email):
        out["reason"] = "personal email address (gmail, yahoo ...) - add it by hand if it's a vendor"
        return out
    if _is_board(email):
        out["reason"] = "a job board's relay address, not the recruiter's own"
        return out
    company, guess, client = resolve_company(email, vendor_company, posting_company, client)
    job_url = job_url if re.match(r"^https?://", job_url or "", re.I) else ""
    cur = conn.cursor()
    row = _find(cur, owner_id, region, email)
    client_note = f"Client: {client}" if client else ""
    if not row:
        cur.execute("""INSERT INTO vendor_contacts (owner_user_id, company_name, company_norm, email, email_domain, contact_name, phone,
                       title, notes, status, source, region, linkedin_url, found_via, last_job_title, last_job_url, last_seen_at, req_count,
                       company_guess) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 'active', 'auto', ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, 1, ?)""",
                    (owner_id, company, normalize_company(company), email, company_domain(email), (name or "")[:200], (phone or "")[:60],
                     (title or "")[:200], client_note, region, (linkedin_url or "")[:300], (found_via or "")[:120], (job_title or "")[:200],
                     job_url[:500], 1 if guess else 0))
        conn.commit()
        row = _find(cur, owner_id, region, email) or {}
        out.update(status="created", id=row.get("id"), company=company, company_guess=guess)
        return out
    fields = {}
    for col, val in (("contact_name", name), ("phone", phone), ("title", title), ("linkedin_url", linkedin_url)):
        if val and not (row.get(col) or "").strip():
            fields[col] = val.strip()[:300]
    if row.get("company_guess") and not guess and company:
        fields["company_name"], fields["company_norm"], fields["company_guess"] = company, normalize_company(company), 0
    notes = row.get("notes") or ""
    if client_note and client.lower() not in notes.lower() and len(notes) < 250:
        fields["notes"] = (notes + ("; " if notes else "") + client_note)[:300]
    if found_via and not row.get("found_via"):
        fields["found_via"] = found_via[:120]
    new_job = bool(job_title or job_url) and (job_url or job_title) != (row.get("last_job_url") or row.get("last_job_title") or "")
    if new_job:
        fields.update(last_job_title=(job_title or "")[:200], last_job_url=job_url[:500], req_count=(row.get("req_count") or 0) + 1)
        if found_via:
            fields["found_via"] = found_via[:120]
    if fields:
        cur.execute(f"UPDATE vendor_contacts SET {', '.join(k + ' = ?' for k in fields)}, last_seen_at = CURRENT_TIMESTAMP WHERE id = ?",
                    list(fields.values()) + [row["id"]])
        conn.commit()
    out.update(status="updated" if fields else "unchanged", id=row["id"], company=fields.get("company_name") or row.get("company_name"),
               company_guess=bool(fields.get("company_guess", row.get("company_guess"))), name=row.get("contact_name") or name or "")
    return out


def capture_from_job(conn, owner_id: int, region: str, job: Dict, email: str = "", vendor_company: str = "", client: str = "") -> Dict:
    """capture() with everything a saved job knows: company, recruiter name / phone, title, link, where it was found;
    a recruiter's LinkedIn post adds their headline and profile link (stored in the job's description)."""
    desc = job.get("description") or ""
    title, linkedin = "", ""
    if (job.get("source") or "") == "Recruiter Post":
        m = re.search(r"^Recruiter post by .+? \((.+)\)\s*$", desc, re.M)
        title = m.group(1) if m else ""
        m = re.search(r"^Profile:\s*(https?://\S+)", desc, re.M)
        linkedin = m.group(1) if m else ""
        # a post has no company field: "Talent Acquisition at CloudHire" names it - when it belongs to the email
        m = re.search(r"\b(?:at|@)\s+([A-Z][\w&.\- ]{1,60}?)(?:\s*[|,(]|$)", title)
        email_ = email or job.get("recruiter_email") or ""
        if m and not vendor_company and company_matches_email(m.group(1), email_):
            vendor_company = m.group(1).strip()
    via = (job.get("source") or "").replace(" (Live 24h)", "").replace(" (Live)", "").strip()
    via = {"Email requirement": "Gmail requirement", "Recruiter Post": "LinkedIn recruiter post"}.get(via, via)
    return capture(conn, owner_id, region, email or job.get("recruiter_email") or "", vendor_company=vendor_company,
                   posting_company=job.get("company") or "", client=client, name=job.get("recruiter_name") or "",
                   phone=job.get("recruiter_phone") or "", title=title, linkedin_url=linkedin, job_title=job.get("title") or "",
                   job_url=job.get("url") or "", found_via=(via + " job") if via and "job" not in via.lower() and "requirement" not in via.lower()
                   and "post" not in via.lower() else via)


def fix_domain_company_names(conn) -> int:
    """Once: contacts saved with only a domain as the company ('pyramid.com', from the old Paste window) get the real
    name from a job with the same recruiter email - when that job's company belongs to the email - else a readable
    domain marked as a guess. Returns how many were changed."""
    cur = conn.cursor()
    cur.execute("SELECT id, email, company_name FROM vendor_contacts")
    changed = 0
    for r in cur.fetchall():
        cid, email, company = r[0], (r[1] or "").lower(), (r[2] or "").strip()
        if not company or "." not in company or company.lower() != company_domain(email):
            continue
        cur.execute("SELECT company FROM jobs WHERE LOWER(recruiter_email) = ? AND company IS NOT NULL ORDER BY id DESC LIMIT 5", (email,))
        names = [x[0] for x in cur.fetchall() if x[0] and company_matches_email(x[0], email)]
        new, guess = (names[0], 0) if names else (pretty_domain(email), 1)
        if new and new != company:
            cur.execute("UPDATE vendor_contacts SET company_name = ?, company_norm = ?, company_guess = ? WHERE id = ?",
                        (new, normalize_company(new), guess, cid))
            changed += 1
    conn.commit()
    return changed
